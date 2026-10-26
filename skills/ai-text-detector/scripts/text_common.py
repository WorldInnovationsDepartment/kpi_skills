"""Shared text helpers for the ai-text-detector and plagiarism-check skills.

Not a standalone script: the other scripts import it from their own directory
(Python puts a script's directory first on sys.path), so it must stay next to
them. Both skills carry an identical copy so that each skill is self-contained;
keep the two copies in sync (tests/*/ compare them).

What it does:
  * extract text from .txt/.md/.tex/.docx/.pdf/.odt/.rtf (the last two via LibreOffice);
  * split it into paragraphs and classify each one (body, heading, reference list,
    quotation, code, formula, table row, caption, front matter);
  * split paragraphs into sentences (Ukrainian + English abbreviations);
  * detect language (uk / ru / en / mixed) from the alphabet;
  * normalise text for matching (case, apostrophes, homoglyphs, whitespace);
  * detect the tampering tricks that KPI's Unicheck instruction forbids
    (Latin letters inside Cyrillic words, invisible characters).

Dependencies beyond the standard library are optional and imported lazily:
python-docx for .docx, PyMuPDF (fitz) for .pdf, LibreOffice for .odt/.rtf/.doc.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED = {".txt", ".md", ".markdown", ".tex", ".docx", ".pdf", ".odt", ".rtf", ".doc", ".html", ".htm"}

CYRILLIC = re.compile(r"[\u0400-\u04FF]")
LATIN = re.compile(r"[A-Za-z]")
UK_ONLY = re.compile(r"[іїєґІЇЄҐ]")
RU_ONLY = re.compile(r"[ыэъёЫЭЪЁ]")
WORD = re.compile(r"[\w'’ʼ-]+", re.UNICODE)
INVISIBLE = re.compile(r"[\u200B-\u200F\u2060\uFEFF\u00AD\u202A-\u202E]")

# Latin -> Cyrillic homoglyphs (the pairs the KPI Unicheck instruction lists, plus a few common ones).
HOMOGLYPHS = str.maketrans({
    "A": "А", "a": "а", "B": "В", "C": "С", "c": "с", "E": "Е", "e": "е", "H": "Н", "I": "І", "i": "і",
    "K": "К", "M": "М", "O": "О", "o": "о", "P": "Р", "p": "р", "T": "Т", "X": "Х", "x": "х", "y": "у",
})

REFERENCE_HEADINGS = re.compile(
    r"^\s*(?:(?:розділ\s+\d+\.?\s*)?(?:список\s+(?:використан(?:их|ої)\s+)?(?:джерел|літератури)|"
    r"перелік\s+(?:використаних\s+)?(?:джерел|посилань)|література|бібліографі(?:я|чний\s+список)|"
    r"джерела|references|bibliography|literature\s+cited|works\s+cited|список\s+джерел))\s*[:.]?\s*$",
    re.I)
APPENDIX_HEADINGS = re.compile(r"^\s*(?:додат(?:ок|ки)|appendix|appendices|annex)\b", re.I)
HEADING_LIKE = re.compile(r"^\s*(?:(?:розділ|глава|chapter|section)\s+\d+|\d+(?:\.\d+)*\.?\s+\S|[A-ZА-ЯІЇЄҐ][^.!?]{0,80})$")
CAPTION = re.compile(r"^\s*(?:рис(?:унок|\.)|табл(?:иця|\.)|fig(?:ure|\.)|table)\s*[\d.]+", re.I)
REFERENCE_ITEM = re.compile(r"^\s*(?:\[\d+\]|\d+[.)])\s+\S")
QUOTE_START = re.compile(r"^\s*[«\"“„]")
FORMULA_CHARS = re.compile(r"[=∑∫√≤≥≈∂∆∇∞→←↔±×÷∈∉∀∃αβγδεζηθλμνξπρστυφχψωΩ^_{}\\]")
SENTENCE_END = re.compile(r"(?<=[.!?…])\s+(?=[«\"“(\[]?[A-ZА-ЯІЇЄҐ0-9])")
ABBREVIATIONS = (
    "т.д", "т.ін", "т.п", "т.ч", "ім", "рис", "табл", "розд", "п", "пп", "ст", "р", "рр", "с", "сс", "вул", "м", "обл",
    "проф", "доц", "акад", "канд", "д-р", "див", "напр", "та ін", "тощо", "№",
    "fig", "figs", "tab", "eq", "eqs", "sec", "ref", "refs", "vol", "no", "pp", "p", "ed", "eds", "et al", "e.g", "i.e",
    "cf", "vs", "approx", "dr", "prof", "mr", "ms", "mrs", "jr", "inc", "ltd", "co", "univ", "dept",
)
_ABBR_RE = re.compile(r"(?:^|\s)(?:" + "|".join(re.escape(a) for a in ABBREVIATIONS) + r")\.$", re.I)


class ExtractError(Exception):
    """The file could not be read or converted to text."""


@dataclass
class Paragraph:
    index: int
    text: str
    kind: str = "body"          # body | heading | reference | quote | code | formula | table | caption | front | appendix
    page: int | None = None

    @property
    def words(self) -> int:
        return len(WORD.findall(self.text))


@dataclass
class Document:
    path: Path
    paragraphs: list[Paragraph]
    source_kind: str
    language: str = "unknown"
    notes: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.paragraphs)

    def body(self) -> list[Paragraph]:
        return [p for p in self.paragraphs if p.kind == "body"]

    @property
    def words(self) -> int:
        return sum(p.words for p in self.paragraphs)


# ----------------------------------------------------------------------------- extraction

def _soffice() -> str | None:
    for cand in ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        found = shutil.which(cand) or (cand if Path(cand).exists() else None)
        if found:
            return found
    return None


def _convert_with_soffice(path: Path, fmt: str = "txt:Text (encoded):UTF8") -> str:
    exe = _soffice()
    if not exe:
        raise ExtractError(f"{path.name}: LibreOffice (soffice) is required to read {path.suffix} files")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([exe, "--headless", "--convert-to", fmt, "--outdir", tmp, str(path)],
                       check=False, capture_output=True, timeout=300)
        outs = list(Path(tmp).glob("*.txt"))
        if not outs:
            raise ExtractError(f"{path.name}: LibreOffice conversion failed")
        return outs[0].read_text(encoding="utf-8", errors="replace")


def _read_docx(path: Path) -> list[str]:
    try:
        from docx import Document as DocxDocument
        from docx.oxml.ns import qn
    except ImportError as exc:  # pragma: no cover
        raise ExtractError("python-docx is required for .docx input") from exc
    doc = DocxDocument(str(path))
    out: list[str] = []
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            text = "".join(t.text or "" for t in child.iter(qn("w:t")))
            if child.findall(".//" + qn("m:oMath")):
                text = (text + " [formula]").strip()
            out.append(text)
        elif child.tag == qn("w:tbl"):
            for row in child.iter(qn("w:tr")):
                cells = []
                for cell in row.iter(qn("w:tc")):
                    cells.append(" ".join((t.text or "") for t in cell.iter(qn("w:t"))).strip())
                out.append("\t".join(cells))
            out.append("")
    return out


def _read_pdf(path: Path) -> list[tuple[str, int]]:
    try:
        import pymupdf as fitz
    except ImportError:  # pragma: no cover
        try:
            import fitz  # older PyMuPDF
        except ImportError as exc:
            raise ExtractError("PyMuPDF is required for .pdf input") from exc
    out: list[tuple[str, int]] = []
    with fitz.open(str(path)) as pdf:
        for pno, page in enumerate(pdf, 1):
            blocks = page.get_text("blocks")
            for b in sorted(blocks, key=lambda b: (round(b[1] / 5), b[0])):
                txt = b[4] if len(b) > 4 else ""
                if not isinstance(txt, str) or not txt.strip():
                    continue
                # Join hyphenated line breaks and hard-wrapped lines inside a block.
                txt = re.sub(r"(\w)-\n(\w)", r"\1\2", txt)
                txt = re.sub(r"\s*\n\s*", " ", txt).strip()
                out.append((txt, pno))
    if not any(t.strip() for t, _ in out):
        raise ExtractError(f"{path.name}: the PDF has no text layer (scanned?). OCR it first (e.g. ocrmypdf).")
    return out


_TEX_DROP_ENV = re.compile(r"\\begin\{(equation\*?|align\*?|gather\*?|figure\*?|table\*?|tabular|tikzpicture|lstlisting|"
                           r"verbatim|thebibliography|algorithm|algorithmic|eqnarray\*?|multline\*?|displaymath)\}.*?"
                           r"\\end\{\1\}", re.S)
_TEX_KEEP_ARG = re.compile(r"\\(?:textbf|textit|emph|underline|texttt|textsc|section\*?|subsection\*?|subsubsection\*?|"
                           r"chapter\*?|paragraph|caption|title|author|footnote)\{([^{}]*)\}")
_TEX_DROP_CMD = re.compile(r"\\(?:cite[pt]?\*?|ref|eqref|label|autoref|cref|Cref|index|vspace\*?|hspace\*?|includegraphics|"
                           r"usepackage|documentclass|bibliography|bibliographystyle|newcommand|renewcommand|input|include)"
                           r"(?:\[[^\]]*\])?\{[^{}]*\}")


def _read_tex(text: str) -> str:
    text = re.sub(r"(?<!\\)%.*", "", text)
    text = _TEX_DROP_ENV.sub("\n[formula]\n", text)
    text = re.sub(r"\$\$.*?\$\$", " [formula] ", text, flags=re.S)
    text = re.sub(r"\\\[.*?\\\]", " [formula] ", text, flags=re.S)
    text = re.sub(r"\$[^$]*\$", " [formula] ", text)
    text = _TEX_DROP_CMD.sub(" ", text)
    for _ in range(3):
        text = _TEX_KEEP_ARG.sub(r"\1", text)
    text = re.sub(r"\\begin\{[^}]*\}|\\end\{[^}]*\}", "\n", text)
    text = re.sub(r"\\item\b", "\n", text)
    text = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?", " ", text)
    text = text.replace("~", " ").replace("{", "").replace("}", "")
    return text


def _read_md(text: str) -> str:
    text = re.sub(r"^---\n.*?\n---\n", "", text, count=1, flags=re.S)        # YAML front matter
    text = re.sub(r"```.*?```", "\n[code]\n", text, flags=re.S)
    text = re.sub(r"\$\$.*?\$\$", " [formula] ", text, flags=re.S)
    text = re.sub(r"\$[^$\n]+\$", " [formula] ", text)
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", " ", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"[*_]{1,3}(\S[^*_]*?)[*_]{1,3}", r"\1", text)
    text = re.sub(r"^\s*[-*+]\s+", "", text, flags=re.M)
    text = re.sub(r"^\s*>\s?", "", text, flags=re.M)
    return text


def _read_html(text: str) -> str:
    text = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)
    text = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</h\d>|</li>|</tr>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    import html as _html
    return _html.unescape(text)


def extract(path: Path | str) -> Document:
    """Read a file into a Document with classified paragraphs."""
    path = Path(path)
    if not path.exists():
        raise ExtractError(f"{path}: file not found")
    suffix = path.suffix.lower()
    notes: list[str] = []
    raw_paras: list[tuple[str, int | None]] = []
    if suffix == ".docx":
        raw_paras = [(t, None) for t in _read_docx(path)]
        kind = "docx"
    elif suffix == ".pdf":
        raw_paras = _read_pdf(path)
        kind = "pdf"
        notes.append("PDF input: paragraph boundaries are reconstructed from text blocks and may be imperfect.")
    elif suffix in {".odt", ".rtf", ".doc"}:
        raw_paras = [(t, None) for t in _convert_with_soffice(path).split("\n")]
        kind = suffix[1:]
    elif suffix == ".tex":
        raw_paras = [(t, None) for t in _read_tex(path.read_text(encoding="utf-8", errors="replace")).split("\n")]
        kind = "tex"
    elif suffix in {".md", ".markdown"}:
        raw_paras = [(t, None) for t in _read_md(path.read_text(encoding="utf-8", errors="replace")).split("\n")]
        kind = "md"
    elif suffix in {".html", ".htm"}:
        raw_paras = [(t, None) for t in _read_html(path.read_text(encoding="utf-8", errors="replace")).split("\n")]
        kind = "html"
    elif suffix in {".txt", ""} or suffix not in SUPPORTED:
        raw_paras = [(t, None) for t in path.read_text(encoding="utf-8", errors="replace").split("\n")]
        kind = "txt"
    else:  # pragma: no cover
        raise ExtractError(f"{path}: unsupported file type {suffix}")
    paragraphs = _build_paragraphs(raw_paras, pre_split=kind in {"docx", "pdf", "odt", "rtf", "doc", "html"})
    doc = Document(path=path, paragraphs=paragraphs, source_kind=kind, notes=notes)
    classify(doc)
    doc.language = detect_language(" ".join(p.text for p in doc.body()) or doc.text)
    return doc


def _build_paragraphs(raw: list[tuple[str, int | None]], pre_split: bool) -> list[Paragraph]:
    """Turn raw entries into paragraphs.

    pre_split=True (docx, pdf, LibreOffice export, html): every non-empty entry is one paragraph already.
    pre_split=False (txt, md, tex): hard-wrapped lines are merged; a blank line, a list item or a
    heading-like short line starts a new paragraph.
    """
    paras: list[Paragraph] = []
    if pre_split:
        for text, pg in raw:
            text = re.sub(r"[ \t]+", " ", text).strip()
            if text:
                paras.append(Paragraph(len(paras), text, page=pg))
        return paras
    buf: list[str] = []

    def flush():
        nonlocal buf
        if buf:
            text = re.sub(r"[ \t]+", " ", " ".join(buf)).strip()
            if text:
                paras.append(Paragraph(len(paras), text))
        buf = []

    for text, _ in raw:
        stripped = text.strip()
        if not stripped:
            flush()
            continue
        list_item = re.match(r"^(?:[-*•]|\d+[.)]|\[\d+\])\s", stripped)
        heading_like = len(stripped) < 60 and not stripped.endswith((",", ";")) and (stripped.isupper() or
                       re.match(r"^(?:розділ|глава|chapter|section|\d+(?:\.\d+)*\.?)\s", stripped, re.I))
        if buf and (list_item or heading_like):
            flush()
        buf.append(stripped)
        if heading_like:
            flush()
    flush()
    return paras


# ----------------------------------------------------------------------------- classification

def classify(doc: Document) -> None:
    """Assign Paragraph.kind in place."""
    in_refs = False
    in_appendix = False
    seen_body = False
    for p in doc.paragraphs:
        t = p.text
        if REFERENCE_HEADINGS.match(t):
            in_refs, in_appendix = True, False
            p.kind = "heading"
            continue
        if APPENDIX_HEADINGS.match(t) and len(t) < 120:
            in_refs, in_appendix = False, True
            p.kind = "heading"
            continue
        if in_refs:
            # The reference list ends at the next section-like heading that is not a numbered reference.
            if (HEADING_LIKE.match(t) and not REFERENCE_ITEM.match(t) and len(t) < 80 and p.words <= 10
                    and not re.search(r"\d{4}", t)):
                in_refs = False
                p.kind = "heading"
                continue
            p.kind = "reference"
            continue
        if t in ("[code]", "[formula]") or (len(FORMULA_CHARS.findall(t)) >= 3 and p.words < 25):
            p.kind = "formula" if "[code]" not in t else "code"
        elif "\t" in t and p.words < 60:
            p.kind = "table"
        elif CAPTION.match(t):
            p.kind = "caption"
        elif QUOTE_START.match(t) and re.search(r"[»\"”]\s*[.,;]?\s*(?:\[\d+[^\]]*\]|\([^)]*\d{4}[^)]*\))?\s*\.?$", t):
            p.kind = "quote"
        elif p.words <= 12 and not t.endswith((".", ",", ";")) and (HEADING_LIKE.match(t) or t.isupper()):
            p.kind = "heading"
        elif not seen_body and p.words < 8 and not t.endswith("."):
            p.kind = "front"      # title page, affiliations, etc. before the first real paragraph
        else:
            p.kind = "appendix" if in_appendix else "body"
            if p.kind == "body" and p.words >= 15:
                seen_body = True
    # An "appendix" that is prose counts as body for AI detection but is reported separately.


# ----------------------------------------------------------------------------- sentences & words

def sentences(text: str) -> list[str]:
    """Split a paragraph into sentences; keeps abbreviations like "рис." and "et al." together."""
    parts = SENTENCE_END.split(text.strip())
    out: list[str] = []
    for part in parts:
        if out and (_ABBR_RE.search(out[-1]) or len(out[-1]) < 8):
            out[-1] = out[-1] + " " + part
        else:
            out.append(part)
    return [s.strip() for s in out if s.strip()]


def words(text: str) -> list[str]:
    return WORD.findall(text)


def word_spans(text: str) -> list[tuple[str, int, int]]:
    return [(m.group(0), m.start(), m.end()) for m in WORD.finditer(text)]


# ----------------------------------------------------------------------------- language & normalisation

def detect_language(text: str) -> str:
    sample = text[:20000]
    cyr, lat = len(CYRILLIC.findall(sample)), len(LATIN.findall(sample))
    if cyr + lat < 20:
        return "unknown"
    if cyr > lat * 2:
        uk, ru = len(UK_ONLY.findall(sample)), len(RU_ONLY.findall(sample))
        if ru > uk * 2:
            return "ru"
        if uk > ru:
            return "uk"
        return "cyrillic"
    if lat > cyr * 2:
        return "en"
    return "mixed"


def normalize(text: str) -> str:
    """Lower-case, unify apostrophes/quotes/dashes, fix homoglyphs inside Cyrillic words, collapse spaces."""
    text = unicodedata.normalize("NFKC", text)
    text = INVISIBLE.sub("", text)
    text = text.replace("’", "'").replace("ʼ", "'").replace("`", "'").replace("‘", "'")
    text = re.sub(r"[«»“”„]", '"', text)
    text = re.sub(r"[‐‑‒–—―]", "-", text)
    text = fix_homoglyphs(text)
    text = text.lower().replace("ё", "е")
    return re.sub(r"\s+", " ", text).strip()


def fix_homoglyphs(text: str) -> str:
    """Replace Latin look-alikes inside words that are otherwise Cyrillic."""
    def repl(m: re.Match) -> str:
        w = m.group(0)
        if CYRILLIC.search(w) and LATIN.search(w):
            if len(CYRILLIC.findall(w)) >= len(LATIN.findall(w)):
                return w.translate(HOMOGLYPHS)
        return w
    return WORD.sub(repl, text)


def tampering_report(text: str) -> dict:
    """Detect the tricks forbidden by the KPI Unicheck instruction (Додаток 1 до Положення)."""
    mixed = [w for w in WORD.findall(text) if CYRILLIC.search(w) and LATIN.search(w)
             and not re.search(r"\d", w) and len(w) > 2]
    invisible = len(INVISIBLE.findall(text))
    return {
        "mixed_script_words": len(mixed),
        "mixed_script_examples": sorted(set(mixed))[:20],
        "invisible_characters": invisible,
        "suspicious": len(mixed) > 5 or invisible > 20,
    }


# ----------------------------------------------------------------------------- chunking for LLM scoring

def chunks_by_words(paragraphs: list[Paragraph], target_words: int = 220, min_words: int = 40) -> list[dict]:
    """Group consecutive body paragraphs into chunks of about target_words words.

    A long paragraph is split at sentence boundaries; short consecutive paragraphs are merged.
    Each chunk records which paragraph indexes it covers.
    """
    out: list[dict] = []
    buf: list[str] = []
    buf_paras: list[int] = []
    buf_words = 0

    def flush():
        nonlocal buf, buf_paras, buf_words
        if buf and buf_words >= min_words:
            out.append({"text": " ".join(buf), "paragraphs": sorted(set(buf_paras)), "words": buf_words})
        elif buf and out:               # too short: append to the previous chunk
            out[-1]["text"] += " " + " ".join(buf)
            out[-1]["paragraphs"] = sorted(set(out[-1]["paragraphs"] + buf_paras))
            out[-1]["words"] += buf_words
        buf, buf_paras, buf_words = [], [], 0

    for p in paragraphs:
        for s in sentences(p.text):
            n = len(words(s))
            if buf_words + n > target_words * 1.4 and buf_words >= min_words:
                flush()
            buf.append(s)
            buf_paras.append(p.index)
            buf_words += n
            if buf_words >= target_words:
                flush()
    flush()
    return out
