"""deck_common.py -- shared loading/parsing for the phd-seminar-review scripts.

Loads a presentation (.pptx, .ppt, .odp, .key via LibreOffice, or .pdf) into a
plain dict:

    {"source": str, "format": "pptx"|"pdf", "converted_from": str|None,
     "slide_width_in": float, "slide_height_in": float, "slide_count": int,
     "slides": [{"index": 1, "title": str, "paragraphs": [str], "text": str,
                 "notes": str, "pictures": int, "tables": int, "charts": int, "shapes": int,
                 "words": int, "font_sizes": [float], "has_slide_number": bool}]}

and finds the parts of a PhD seminar talk (title slide, relevance, object,
subject, aim, tasks, novelty, conclusions, publications, ...) in it.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import unicodedata
import urllib.parse
from pathlib import Path

CONVERTIBLE = {".ppt", ".pps", ".pot", ".odp", ".key", ".pptm", ".ppsx", ".potx"}
EMU_PER_IN = 914400


class DeckError(Exception):
    pass


# ---------------------------------------------------------------- text utils

APOS = "'’ʼ`´"


# Ukrainian -> Latin, the official passport transliteration (КМУ №55, 27.01.2010).
_TRANSLIT = dict(zip("абвгґдезиіїйклмнопрстуфхцчшщюяєжь'", [
    "a", "b", "v", "h", "g", "d", "e", "z", "y", "i", "i", "i", "k", "l", "m", "n", "o", "p", "r", "s", "t", "u",
    "f", "kh", "ts", "ch", "sh", "shch", "iu", "ia", "ie", "zh", "", ""]))


def translit(s: str) -> str:
    """Latin spelling of a Ukrainian name, e.g. «Петренко» -> «petrenko», «Барабаш» -> «barabash»."""
    return "".join(_TRANSLIT.get(c, c) for c in norm(s))


def norm(s: str) -> str:
    """Lower-case, unify apostrophes/quotes/dashes, squeeze whitespace."""
    s = unicodedata.normalize("NFC", s or "")
    for a in APOS:
        s = s.replace(a, "'")
    s = re.sub(r"[«»“”„\"]", '"', s)
    s = re.sub(r"[–—−]", "-", s)
    return re.sub(r"\s+", " ", s).strip().lower()


def words(s: str) -> list[str]:
    return re.findall(r"[\w'’ʼ-]+", s or "")


def is_mostly_upper(s: str) -> bool:
    letters = [c for c in s if c.isalpha()]
    return bool(letters) and sum(c.isupper() for c in letters) / len(letters) > 0.7


# ---------------------------------------------------------------- conversion

def find_soffice() -> str | None:
    for c in ("soffice", "libreoffice", "/opt/homebrew/bin/soffice", "/usr/local/bin/soffice",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if shutil.which(c) or Path(c).is_file():
            return c
    return None


def soffice_convert(src: Path, fmt: str, out_dir: Path) -> Path:
    """Convert `src` to `fmt` ("pptx" or "pdf") with LibreOffice into `out_dir`."""
    soffice = find_soffice()
    if not soffice:
        raise DeckError(f"{src.name}: converting {src.suffix} needs LibreOffice (soffice), which was not found. "
                        "Install LibreOffice or save the deck as .pptx/.pdf.")
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as profile:
        # A private profile keeps parallel soffice runs from fighting over the lock file.
        cmd = [soffice, f"-env:UserInstallation=file://{urllib.parse.quote(profile)}", "--headless",
               "--convert-to", fmt, "--outdir", str(out_dir), str(src)]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    out = out_dir / (src.stem + "." + fmt)
    if r.returncode != 0 or not out.is_file():
        raise DeckError(f"{src.name}: LibreOffice could not convert to {fmt}: {(r.stderr or r.stdout).strip()[:300]}")
    return out


# ---------------------------------------------------------------- loaders

def load_deck(path: Path, work_dir: Path | None = None) -> dict:
    path = Path(path)
    if not path.is_file():
        raise DeckError(f"{path}: no such file")
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _load_pdf(path)
    if ext in (".pptx",):
        return _load_pptx(path)
    if ext in CONVERTIBLE:
        tmp = Path(work_dir) if work_dir else Path(tempfile.mkdtemp(prefix="seminar-deck-"))
        deck = _load_pptx(soffice_convert(path, "pptx", tmp))
        deck["source"], deck["converted_from"] = str(path), ext.lstrip(".")
        return deck
    raise DeckError(f"{path.name}: unsupported format {ext!r} (use .pptx, .ppt, .odp, .key or .pdf)")


def _shape_texts(shape, out: list, sizes: list, counts: dict, slide_h: int, top_offset: int = 0):
    """Collect (top, size, text) for every text frame, recursing into groups."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            _shape_texts(sub, out, sizes, counts, slide_h)
        return
    if st == MSO_SHAPE_TYPE.PICTURE:
        counts["pictures"] += 1
    elif st in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.LINE):
        counts["shapes"] += 1
    if getattr(shape, "has_chart", False) and shape.has_chart:
        counts["charts"] += 1
    if getattr(shape, "has_table", False) and shape.has_table:
        counts["tables"] += 1
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in shape.table.rows]
        out.append({"top": shape.top or 0, "size": None, "paras": [r for r in rows if r.strip(" |")],
                    "title_ph": False})
        return
    if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
        return
    paras, max_size = [], None
    for p in shape.text_frame.paragraphs:
        t = "".join(r.text for r in p.runs) or p.text
        t = t.replace("\x0b", " ").strip()
        for r in p.runs:
            if r.font.size is not None and r.text.strip():
                sz = r.font.size.pt
                sizes.append(sz)
                max_size = sz if max_size is None else max(max_size, sz)
        if t:
            paras.append(t)
    xml = shape._element.xml
    if 'type="slidenum"' in xml:
        counts["slidenum"] = True
        return
    if not paras:
        return
    is_title = False
    if shape.is_placeholder:
        try:
            is_title = shape.placeholder_format.type in (1, 3)  # TITLE, CENTER_TITLE
        except ValueError:
            pass
    out.append({"top": shape.top or 0, "size": max_size, "paras": paras, "title_ph": is_title})


def _pick_title(blocks: list, slide_h: int) -> str:
    for b in blocks:
        if b["title_ph"]:
            return " ".join(b["paras"])
    top = [b for b in blocks if b["top"] < slide_h * 0.22 and not re.fullmatch(r"\d{1,3}", " ".join(b["paras"]))]
    if not top:
        return ""
    # Prefer a real heading over short label boxes ("НР.1"), then all caps, the largest font, the highest box.
    def is_label(b):
        t = " ".join(b["paras"])
        return len(t) < 12 and bool(re.search(r"\d", t))

    top.sort(key=lambda b: (is_label(b), not is_mostly_upper(" ".join(b["paras"])),
                            -(b["size"] or 0), b["top"]))
    return " ".join(top[0]["paras"])


def _load_pptx(path: Path) -> dict:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as e:  # noqa: BLE001 - python-pptx raises many kinds
        raise DeckError(f"{path.name}: cannot open as .pptx ({e})") from e
    sw, sh = prs.slide_width or 12192000, prs.slide_height or 6858000
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        blocks, sizes = [], []
        counts = {"pictures": 0, "tables": 0, "charts": 0, "shapes": 0, "slidenum": False}
        for shape in slide.shapes:
            _shape_texts(shape, blocks, sizes, counts, sh)
        blocks.sort(key=lambda b: b["top"])
        title = _pick_title(blocks, sh)
        paras = [p for b in blocks for p in b["paras"]]
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        text = "\n".join(paras)
        slides.append({
            "index": i, "title": title, "paragraphs": paras, "text": text, "notes": notes,
            "pictures": counts["pictures"], "tables": counts["tables"], "charts": counts["charts"],
            "shapes": counts["shapes"], "words": len(words(text)), "font_sizes": sorted(set(sizes)),
            "has_slide_number": counts["slidenum"] or _has_layout_slidenum(slide)
            or any(b["paras"] == [str(i)] for b in blocks),
        })
    return {"source": str(path), "format": "pptx", "converted_from": None,
            "slide_width_in": round(sw / EMU_PER_IN, 2), "slide_height_in": round(sh / EMU_PER_IN, 2),
            "slide_count": len(slides), "slides": slides}


def _has_layout_slidenum(slide) -> bool:
    # A slide-number placeholder only shows when the slide itself carries it.
    for shape in slide.placeholders:
        try:
            if shape.placeholder_format.type == 13:  # SLIDE_NUMBER
                return True
        except ValueError:
            pass
    return False


def _load_pdf(path: Path) -> dict:
    try:
        import pymupdf as fitz
    except ImportError:  # pragma: no cover
        import fitz
    doc = fitz.open(str(path))
    slides = []
    w = h = 0
    for i, page in enumerate(doc, 1):
        w, h = page.rect.width, page.rect.height
        lines, sizes = [], []
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                spans = [s for s in line["spans"] if s["text"].strip()]
                if not spans:
                    continue
                t = " ".join(s["text"].strip() for s in spans)
                sz = max(s["size"] for s in spans)
                sizes += [round(s["size"], 1) for s in spans]
                lines.append({"top": line["bbox"][1], "size": sz, "text": t})
        lines.sort(key=lambda l: l["top"])
        top = [l for l in lines if l["top"] < h * 0.22 and not re.fullmatch(r"\d{1,3}", l["text"])]
        title = ""
        if top:
            # Same preference as for .pptx: an all-caps heading first, then the largest font.
            caps = [l for l in top if is_mostly_upper(l["text"])] or top
            big = max(l["size"] for l in caps)
            title = " ".join(l["text"] for l in caps if abs(l["size"] - big) < 0.6)
        paras = _join_pdf_lines(lines)
        text = "\n".join(paras)
        last = lines[-1]["text"] if lines else ""
        slides.append({
            "index": i, "title": title, "paragraphs": paras, "text": text, "notes": "",
            "pictures": len(page.get_images()), "tables": 0, "charts": 0,
            "shapes": len(page.get_drawings()),
            "words": len(words(text)), "font_sizes": sorted(set(sizes)),
            "has_slide_number": any(re.fullmatch(rf"{i}(\s*/\s*\d+)?", l["text"]) for l in lines),
        })
    return {"source": str(path), "format": "pdf", "converted_from": None,
            "slide_width_in": round(w / 72, 2), "slide_height_in": round(h / 72, 2),
            "slide_count": len(slides), "slides": slides}


def _join_pdf_lines(lines: list) -> list[str]:
    """Glue wrapped PDF lines back into paragraphs (new paragraph on numbering or a gap)."""
    paras, prev = [], None
    for l in lines:
        t = l["text"]
        start = re.match(r"^(\d{1,2}[.)]\s|[•▪●\-–]\s)", t)
        gap = prev is not None and (l["top"] - prev["top"]) > prev["size"] * 1.9
        if not paras or start or gap or abs(l["size"] - prev["size"]) > 1.5:
            paras.append(t)
        else:
            paras[-1] = paras[-1].rstrip("-") + ("" if paras[-1].endswith("-") else " ") + t
        prev = l
    return paras


# ---------------------------------------------------------------- dissertation text

def load_text(path: Path) -> str:
    """Plain text of the dissertation / autoabstract (.docx, .pdf, .md, .txt, .doc/.odt/.rtf via soffice)."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        try:
            import pymupdf as fitz
        except ImportError:  # pragma: no cover
            import fitz
        return "\n".join(p.get_text() for p in fitz.open(str(path)))
    if ext in (".doc", ".odt", ".rtf"):
        path = soffice_convert(path, "docx", Path(tempfile.mkdtemp(prefix="seminar-diss-")))
        ext = ".docx"
    if ext == ".docx":
        from docx import Document

        d = Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text for c in r.cells))
        return "\n".join(parts)
    raise DeckError(f"{path.name}: unsupported dissertation format {ext!r} (use .docx, .pdf, .md, .txt)")


# ---------------------------------------------------------------- sections

# Each section: regex on the normalised slide title, regex for a labelled
# statement inside body text ("Об'єкт дослідження: ..."), human label.
SECTIONS = {
    "relevance":   (r"актуальн|relevance|motivation", None, "Актуальність теми"),
    "connection":  (r"зв'язок .*(програм|план|тем)", r"зв'язок роботи з науков", "Зв'язок з науковими програмами, планами, темами"),
    "object":      (r"об'єкт", r"об'єкт\w*(\s+дослідж\w*)?\s*[:–—-]", "Об'єкт дослідження"),
    "subject":     (r"предмет", r"предмет\w*(\s+дослідж\w*)?\s*[:–—-]", "Предмет дослідження"),
    "aim":         (r"\bмета\b|\bціль\b|\baim\b|\bgoal\b", r"\bмета\b(\s+\w+){0,3}\s*[:–—-]", "Мета дослідження"),
    "tasks":       (r"завдан|\btasks?\b|objectives", r"(для досягнення мети|поставлено|вирішено|розв'язано)\s.{0,40}завдан", "Завдання дослідження"),
    "analysis":    (r"аналіз|огляд|стан питання|state of the art|related work", None, "Аналіз стану питання / публікацій"),
    "methods":     (r"метод\w*\s+дослідж|methods", r"метод\w*\s+дослідж\w*\s*[:–—-]", "Методи дослідження"),
    "novelty":     (r"новизн|наукові результати|novelty|contribution", r"наукова новизна", "Наукова новизна"),
    "practical":   (r"практичн|впровадж|реаліза|practical", r"практичн\w+\s+значенн", "Практичне значення / впровадження"),
    "contribution": (r"особист\w+ внес", r"особист\w+ внес", "Особистий внесок здобувача"),
    "approbation": (r"апробац|конференц|approbation", r"апробац", "Апробація результатів"),
    "publications": (r"публікац|publications|список праць", r"(статт|публікаці)", "Публікації"),
    "conclusions": (r"висновк|conclusion", None, "Висновки"),
    "thanks":      (r"дякую|thank", None, "Дякую за увагу"),
}

# A title like "Аналіз останніх публікацій" is the literature review, not the list of own papers.
NOT_PUBLICATIONS = r"аналіз|огляд|останніх"


def classify_slide(slide: dict) -> list[str]:
    t = norm(slide["title"])
    found = []
    for key, (title_re, _, _) in SECTIONS.items():
        if title_re and re.search(title_re, t):
            if key == "publications" and re.search(NOT_PUBLICATIONS, t):
                continue
            if key == "methods" and "новизн" in t:
                continue
            found.append(key)
    return found


def section_map(deck: dict) -> dict[str, list[int]]:
    """section -> list of slide indexes (titles first, then labelled body statements)."""
    m: dict[str, list[int]] = {k: [] for k in SECTIONS}
    for s in deck["slides"]:
        for k in classify_slide(s):
            m[k].append(s["index"])
    for s in deck["slides"]:
        body = norm(s["text"])
        for k, (_, label_re, _) in SECTIONS.items():
            if label_re and not m[k] and re.search(label_re, body):
                if k in ("publications", "tasks"):
                    continue  # too generic in body text; only trust titles for these
                m[k].append(s["index"])
    last = deck["slides"][-1] if deck["slides"] else None
    if last and not m["thanks"] and re.search(SECTIONS["thanks"][0], norm(last["text"])) and last["words"] < 25:
        m["thanks"].append(last["index"])
    return m


def labelled_statement(deck: dict, key: str) -> tuple[int, str] | None:
    """Text of "Об'єкт дослідження: ..." style statements (object/subject/aim/methods)."""
    pats = {
        "object": r"об['’ʼ]єкт\w*(?:\s+дослідж\w*)?\s*[:–—-]\s*",
        "subject": r"предмет\w*(?:\s+дослідж\w*)?\s*[:–—-]\s*",
        "aim": r"мета(?:\s+[\w']+){0,3}\s*[:–—-]\s*",
        "methods": r"метод\w*\s+дослідж\w*\s*[:–—-]\s*",
    }
    rx = re.compile(r"^\s*" + pats[key], re.I)
    for s in deck["slides"]:
        for p in s["paragraphs"]:
            mt = rx.search(p)
            if mt and len(p) - mt.end() > 10:
                return s["index"], p[mt.end():].strip()
    # Fallback: a slide titled just "Об'єкт дослідження" -> its first body paragraph.
    for s in deck["slides"]:
        if key in classify_slide(s) and len(classify_slide(s)) == 1:
            body = [p for p in s["paragraphs"] if norm(p) != norm(s["title"]) and len(p) > 15]
            if body:
                return s["index"], body[0]
    return None


ITEM_START = re.compile(r"^\s*(\d{1,2})\s*[.)]\s*|^\s*[•▪●■◆➢✓\-–]\s+")


def numbered_items(deck: dict, key: str, min_len: int = 25) -> list[tuple[int, str]]:
    """Items (numbered paragraphs, or one per paragraph) on the slides of a section."""
    sm = section_map(deck)
    idx = [i for i in sm.get(key, []) if key in classify_slide(deck["slides"][i - 1])]
    items: list[tuple[int, str]] = []
    for i in idx:
        s = deck["slides"][i - 1]
        cur = None
        for p in s["paragraphs"]:
            if norm(p) == norm(s["title"]) or re.fullmatch(r"\d{1,3}", p.strip()):
                continue
            if ITEM_START.match(p) or cur is None:
                if cur:
                    items.append(cur)
                cur = (i, ITEM_START.sub("", p).strip())
            else:
                cur = (cur[0], cur[1] + " " + p.strip())
        if cur:
            items.append(cur)
    return [it for it in items if len(it[1]) >= min_len]
