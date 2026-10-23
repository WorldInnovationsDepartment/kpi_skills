# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "python-docx>=1.1",
#     "pyyaml>=6",
#     "pymupdf>=1.24",
# ]
# ///
"""check_paper.py -- check a manuscript against the requirements of
«Адаптивні системи автоматичного управління» (https://asac.kpi.ua).

    uv run check_paper.py paper.md                       # content checks on the source
    uv run check_paper.py paper.md --docx paper.docx     # + layout checks on the built file
    uv run check_paper.py paper.md --docx paper.docx --pdf preview/paper.pdf   # + page count

Prints one line per finding:

    ERROR <code>: ...   violates a stated requirement -> fix before submitting
    WARN  <code>: ...   a recommendation of the author guide, or a likely problem
    INFO  <code>: ...   reminders that cannot be checked automatically

then `SUMMARY: E error(s), W warning(s), I info`, and exits 1 if there is any
ERROR (0 otherwise). Every rule and its source is listed in
references/requirements.md.
"""
from __future__ import annotations

import argparse
import datetime
import re
import sys
from pathlib import Path

from asac_common import (
    ABSTRACT_PARTS,
    FONT,
    FONT_SIZE_PT,
    FUNDING_PREFIX,
    LANGS,
    LINE_SPACING,
    MARGIN_MM,
    PAGE_H_MM,
    PAGE_W_MM,
    DISPLAY_MATH_RE,
    IMAGE_RE,
    SourceError,
    find_citations,
    has_doi,
    load_source,
    reference_year,
    required_errors,
    section_kind,
    sections,
    strip_code,
    strip_markup,
    table_blocks,
    text_of,
)

ORCID_RE = re.compile(r"^\d{4}-\d{4}-\d{4}-\d{3}[\dX]$")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
COMMON_ABBR = {"IT", "ІТ", "AI", "ШІ", "IOT", "GPS", "API", "SQL", "UAV", "БПЛА", "ПЗ", "CPU", "GPU", "PID", "ПІД",
               "HTTP", "XML", "JSON", "КПІ", "NTUU", "НТУУ", "USA", "EU", "ЄС", "UDC", "УДК", "DOI", "ORCID"}


class Report:
    def __init__(self):
        self.items: list[tuple[str, str, str]] = []

    def add(self, level: str, code: str, msg: str) -> None:
        self.items.append((level, code, msg))

    def error(self, code, msg):
        self.add("ERROR", code, msg)

    def warn(self, code, msg):
        self.add("WARN", code, msg)

    def info(self, code, msg):
        self.add("INFO", code, msg)

    def count(self, level):
        return sum(1 for lv, _, _ in self.items if lv == level)

    def print(self) -> None:
        order = {"ERROR": 0, "WARN": 1, "INFO": 2}
        for lv, code, msg in sorted(self.items, key=lambda x: order[x[0]]):
            print(f"{lv:<5} {code}: {msg}")
        print(f"SUMMARY: {self.count('ERROR')} error(s), {self.count('WARN')} warning(s), {self.count('INFO')} info")


def words(text: str) -> int:
    return len([w for w in re.findall(r"\S+", text) if re.search(r"[\w\d]", w)])


def sentences(text: str) -> list[str]:
    text = re.sub(r"\$[^$]*\$", "x", text)
    text = re.sub(r"\b(e\.g|i\.e|et al|Fig|Tab|Eq|рис|табл|див|т\.д|т\.п|напр|No|Vol|P|С|Т|№)\.", r"\1", text,
                  flags=re.I)
    return [s for s in re.split(r"(?<=[.!?…])\s+(?=[A-ZА-ЯІЇЄҐ0-9«\"(])", text.strip()) if words(s) >= 3]


def abbreviations(text: str) -> list[str]:
    toks = set(re.findall(r"\b[A-ZА-ЯІЇЄҐ][A-ZА-ЯІЇЄҐ0-9]{1,}[a-zа-яіїєґ]?\b", text))
    return sorted(t for t in toks if t.upper() not in COMMON_ABBR and not t.isdigit())


# ---------------------------------------------------------------- content checks (source)

def find_todos(value, path: str = "") -> list[str]:
    if isinstance(value, dict):
        return [p for k, v in value.items() for p in find_todos(v, f"{path}.{k}" if path else str(k))]
    if isinstance(value, list):
        return [p for i, v in enumerate(value) for p in find_todos(v, f"{path}[{i}]")]
    return [path] if isinstance(value, str) and re.search(r"\bTODO\b", value) else []


def check_meta(meta: dict, rep: Report, year_now: int) -> None:
    for e in required_errors(meta):
        rep.error("meta.required", e)
    todos = find_todos(meta)
    if todos:
        rep.error("meta.todo", f"TODO placeholder(s) left in: {', '.join(todos)}")
    lang = meta.get("language") if meta.get("language") in LANGS else "en"

    # Title
    for lg in LANGS:
        t = strip_markup(str(meta.get(f"title_{lg}") or ""))
        if not t:
            continue
        if words(t) > 20:
            rep.warn(f"title.{lg}.length", f"title has {words(t)} words; the guide asks for a short, informative title")
        if t.isupper():
            rep.warn(f"title.{lg}.case", "write the title in normal case; the builder upper-cases it in the title block")
        else:
            abbr = abbreviations(t)
            if abbr:
                rep.warn(f"title.{lg}.abbreviations",
                         f"title contains abbreviation(s) {abbr}; the guide forbids atypical acronyms in the title")

    # Abstracts
    for lg in LANGS:
        v = meta.get(f"abstract_{lg}")
        txt = text_of(v)
        if not txt:
            continue
        n_words, n_chars = words(txt), len(txt)
        if n_chars > 1800:
            rep.error(f"abstract.{lg}.chars", f"abstract_{lg} is {n_chars} characters; the maximum is 1800")
        if not 200 <= n_words <= 250:
            rep.warn(f"abstract.{lg}.words", f"abstract_{lg} has {n_words} words; required 200-250")
        if isinstance(v, dict):
            missing = [names[lg] for key, names in ABSTRACT_PARTS if not str(v.get(key) or "").strip()]
            if missing:
                rep.error(f"abstract.{lg}.structure", f"abstract_{lg} lacks part(s): {', '.join(missing)}")
            purpose = str(v.get("purpose") or "")
            if purpose and len(sentences(purpose)) > 2:
                rep.warn(f"abstract.{lg}.purpose", "the purpose part should be 1-2 sentences")
        else:
            rep.warn(f"abstract.{lg}.structure",
                     f"abstract_{lg} is plain text; the guide requires the structured form "
                     f"({', '.join(n[lg] for _, n in ABSTRACT_PARTS)}) -- give it as a mapping")
        if find_citations(txt):
            rep.error(f"abstract.{lg}.citations", f"abstract_{lg} contains reference citations; they are not allowed")
        abbr = abbreviations(txt)
        if abbr:
            rep.warn(f"abstract.{lg}.abbreviations",
                     f"abstract_{lg} uses abbreviation(s) {abbr}; only the most common ones are allowed")
    top = meta.get("abstract_top")
    if top:
        n = len(text_of(top))
        rep.info("abstract.top", f"abstract_top is {n} characters (the official sample suggests 150-300 characters "
                                 "for the short abstract under the title)")
    else:
        rep.info("abstract.top", f"the title block repeats abstract_{lang} (as in recently published issues); set "
                                 "abstract_top for a short 150-300 character version as in the official sample")

    # Keywords
    for lg in LANGS:
        kws = [strip_markup(str(k)).strip() for k in (meta.get(f"keywords_{lg}") or []) if str(k).strip()]
        if not kws:
            continue
        if len(kws) < 5:
            rep.error(f"keywords.{lg}.count", f"{len(kws)} keywords in keywords_{lg}; at least 5 are required")
        elif len(kws) > 20:
            rep.error(f"keywords.{lg}.count", f"{len(kws)} keywords in keywords_{lg}; at most 20 are allowed")
        elif len(kws) > 10:
            rep.warn(f"keywords.{lg}.count", f"{len(kws)} keywords in keywords_{lg}; the author guide says up to 10")
        single = [k for k in kws if len(k.split()) == 1]
        if single:
            rep.warn(f"keywords.{lg}.single", f"single-word keyword(s) {single} in keywords_{lg}; prefer phrases")
        title_low = strip_markup(str(meta.get(f"title_{lg}") or "")).lower()
        dup = [k for k in kws if k.lower() in title_low]
        if dup:
            rep.info(f"keywords.{lg}.title", f"keyword(s) {dup} repeat the title; keywords should add to it")
    if meta.get("keywords_uk") and meta.get("keywords_en") and \
            len(meta["keywords_uk"]) != len(meta["keywords_en"]):
        rep.warn("keywords.parity", "keywords_uk and keywords_en have different lengths; they should correspond")

    # Authors
    authors = meta.get("authors") or []
    if isinstance(authors, list) and authors:
        corresponding = [a for a in authors if isinstance(a, dict) and a.get("corresponding")]
        if not corresponding:
            rep.error("authors.corresponding", "mark one author with `corresponding: true` (with a phone number)")
        for i, a in enumerate(authors):
            if not isinstance(a, dict):
                continue
            who = a.get("surname_en") or a.get("surname_uk") or f"#{i + 1}"
            orcid = str(a.get("orcid") or "").strip().removeprefix("https://orcid.org/")
            if not orcid:
                rep.error("authors.orcid", f"{who}: ORCID is mandatory")
            elif not ORCID_RE.match(orcid):
                rep.error("authors.orcid", f"{who}: ORCID '{orcid}' is not in the 0000-0000-0000-0000 form")
            elif orcid.startswith("0000-0000-0000-000"):
                rep.warn("authors.placeholder", f"{who}: ORCID {orcid} looks like a placeholder")
            email = str(a.get("email") or "").strip()
            if not email:
                rep.error("authors.email", f"{who}: e-mail is required")
            elif not EMAIL_RE.match(email):
                rep.error("authors.email", f"{who}: e-mail '{email}' is malformed")
            elif email.endswith("@example.com"):
                rep.warn("authors.placeholder", f"{who}: e-mail {email} looks like a placeholder")
            if a.get("corresponding") and not str(a.get("phone") or "").strip():
                rep.error("authors.phone", f"{who}: the corresponding author must give a contact phone")
            if str(a.get("phone") or "").count("0") >= 10:
                rep.warn("authors.placeholder", f"{who}: phone looks like a placeholder")
            for lg in LANGS:
                if not str(a.get(f"position_{lg}") or "").strip():
                    rep.error("authors.position", f"{who}: position_{lg} (place of work and position) is required")
                if not str(a.get(f"degree_{lg}") or "").strip():
                    rep.warn("authors.degree", f"{who}: degree_{lg} is empty (write e.g. 'без ступеня'/'аспірант' "
                                               "if there is none)")
            if a.get("patronymic_uk") in (None, ""):
                rep.warn("authors.patronymic", f"{who}: patronymic_uk is empty (the guide asks for the full name)")
            if not a.get("scopus_id") and not a.get("researcher_id"):
                rep.info("authors.ids", f"{who}: no ScopusID/ResearcherID (only needed if they exist)")

    # Funding / disclosure
    funding = str(meta.get("funding") or "").strip()
    if funding and not funding.startswith(FUNDING_PREFIX[lang]):
        rep.warn("funding.prefix", f"funding must start with \"{FUNDING_PREFIX[lang]}...\"")
    if meta.get("article_type") == "review" and not (funding or meta.get("conflict_of_interest")):
        rep.error("review.disclosure", "review articles must include a disclosure-of-interests statement "
                                       "(funding or conflict_of_interest)")

    # References
    refs = [str(r).strip() for r in (meta.get("references") or []) if str(r).strip()]
    if refs:
        recent = [r for r in refs if (y := reference_year(r)) and year_now - 5 <= y <= year_now]
        if len(refs) < 10:
            rep.warn("refs.count", f"{len(refs)} references; the guide expects a substantive review "
                                   "(10-15 recent international papers recommended)")
        if len(recent) < 10:
            rep.warn("refs.recent", f"{len(recent)} of {len(refs)} references are from {year_now - 5}-{year_now}; "
                                    "the guide recommends 10-15 papers from international journals of the last 5 years")
        no_doi = [i for i, r in enumerate(refs, 1) if not has_doi(r)]
        if no_doi:
            rep.warn("refs.doi", f"no DOI in reference(s) {no_doi}; give the DOI wherever one exists")
        bad = [i for i, r in enumerate(refs, 1) if re.search(r"unpublished|personal communication|неопубл|особист\w* "
                                                               r"(комунікац|повідомл)", r, re.I)]
        if bad:
            rep.error("refs.unpublished", f"reference(s) {bad} cite unpublished data/personal communication; "
                                          "move them into the text, e.g. (Журавель І., 2026, неопубліковані дані)")
        no_year = [i for i, r in enumerate(refs, 1) if reference_year(r) is None]
        if no_year:
            rep.warn("refs.year", f"no publication year found in reference(s) {no_year}")
        undstu = [i for i, r in enumerate(refs, 1) if " – " not in r and " — " not in r and ". " not in r]
        if undstu:
            rep.warn("refs.format", f"reference(s) {undstu} do not look like DSTU 8302:2015 entries")
        numbered = [i for i, r in enumerate(refs, 1) if re.match(r"^\s*\[?\d+[.)\]]\s", r)]
        if numbered:
            rep.error("refs.numbering", f"reference(s) {numbered} start with their own number; the builder numbers "
                                        "the list -- remove it")


def check_body(meta: dict, body: str, rep: Report) -> None:
    lang = meta.get("language") if meta.get("language") in LANGS else "en"
    clean = strip_code(body)
    if re.search(r"\bTODO\b", clean):
        rep.error("body.todo", "TODO placeholder(s) left in the body")
    secs = sections(body, 1)
    kinds = [section_kind(s.title) for s in secs]
    if not secs:
        rep.error("body.sections", "the body has no '# Section' headings")
        return

    if any(re.match(r"(?i)^(references|список|bibliograph|література|abstract|анотація|keywords|ключові)",
                    s.title) for s in secs):
        rep.error("body.generated", "remove Abstract/Keywords/References sections from the body -- they come from the "
                                    "front matter and are generated in the right places")

    atype = meta.get("article_type", "original")
    needed = ["introduction", "conclusions"]
    wanted = ["methods", "results", "discussion"] if atype in ("original", "short") else []
    for k in needed:
        if k not in kinds:
            rep.error(f"body.{k}", f"no '{k.capitalize()}' section (heading containing e.g. "
                                   f"'{'Introduction/Вступ' if k == 'introduction' else 'Conclusions/Висновки'}')")
    for k in wanted:
        if k not in kinds and not (k == "discussion" and any(s and "result" in s for s in kinds)):
            rep.warn(f"body.{k}", f"no '{k}' section; the guide's structure for original papers is Introduction, "
                                  "Materials and methods, Results, Discussion, Conclusions (synonyms are accepted)")
    if kinds and kinds[0] != "introduction":
        rep.warn("body.order", f"the first section is '{secs[0].title}', expected the Introduction")
    if kinds and kinds[-1] != "conclusions":
        rep.warn("body.order", f"the last section is '{secs[-1].title}', expected the Conclusions")

    # Introduction rules
    for s, k in zip(secs, kinds):
        if k == "introduction":
            if IMAGE_RE.search(s.text) or table_blocks(s.text):
                rep.error("intro.floats", "the Introduction must not contain figures or tables")
            n_cites = len({n for _, n in find_citations(s.text)})
            if n_cites < 10:
                rep.warn("intro.citations", f"the Introduction cites {n_cites} distinct source(s); the guide recommends "
                                            "10-15 recent international papers there")
            paras = [p for p in re.split(r"\n\s*\n", s.text) if p.strip() and not p.lstrip().startswith("#")]
            if paras and not re.search(r"(?i)task|aim|purpose|goal|objective|problem|завданн|мет[аою]|задач|проблем",
                                       paras[-1]):
                rep.warn("intro.problem", "the Introduction should end with the problem statement "
                                          "(the specific task/purpose of the work)")
        if k == "conclusions":
            text = re.sub(r"(?m)^#.*$", "", s.text)
            n = len(sentences(text))
            if not 5 <= n <= 10:
                rep.warn("conclusions.length", f"Conclusions have ~{n} sentences; the guide says 5-10")
            if not re.search(r"(?i)future|further|next step|подальш|перспектив|надалі", text):
                rep.warn("conclusions.future", "Conclusions should end with directions for further research")

    # Figures
    figs = list(IMAGE_RE.finditer(clean))
    fig_ref_re = re.compile(r"(?i)(?:fig(?:ure)?s?\.?|рис(?:ун\w*)?\.?)\s*(\d+)")
    for n, m in enumerate(figs, 1):
        if not m.group("alt").strip():
            rep.error("figure.caption", f"figure {n} ({m.group('src')}) has no caption: write ![Caption](file)")
        mentions = [r.start() for r in fig_ref_re.finditer(clean) if int(r.group(1)) == n and not
                    IMAGE_RE.match(clean, max(0, r.start() - 2))]
        if not mentions:
            rep.warn("figure.ref", f"figure {n} is never referred to in the text (e.g. 'Fig. {n}' / 'рис. {n}')")
        elif min(mentions) > m.start():
            rep.warn("figure.ref", f"figure {n} appears before its first mention; place it right after it")
        src = Path(m.group("src"))
        if not src.is_absolute():
            src = SOURCE_DIR / src
        if not src.exists():
            rep.error("figure.file", f"figure {n}: file not found: {m.group('src')}")

    # Tables
    tables = table_blocks(body)
    tab_ref_re = re.compile(r"(?i)(?:tab(?:le)?s?\.?|табл(?:иц\w*)?\.?)\s*(\d+)")
    for n, (pos, cap) in enumerate(tables, 1):
        if not cap:
            rep.error("table.caption", f"table {n} has no caption: add a 'Table: <title>' line before it")
        mentions = [r.start() for r in tab_ref_re.finditer(clean) if int(r.group(1)) == n
                    and not re.match(r"\s*(?:Table)?:", clean[clean.rfind("\n", 0, r.start()) + 1: r.start()] or "x")]
        if not mentions:
            rep.warn("table.ref", f"table {n} is never referred to in the text (e.g. 'Table {n}' / 'табл. {n}')")
        elif min(mentions) > pos:
            rep.warn("table.ref", f"table {n} appears before its first mention; place it right after it")

    # Equations
    if DISPLAY_MATH_RE.search(clean):
        rep.info("equations", f"{len(DISPLAY_MATH_RE.findall(clean))} display equation(s) will be numbered (1), (2)... "
                              "as native Word equations; refer to them by these numbers")
    if re.search(r"!\[[^\]]*\]\([^)]*(formula|equation|eq\d|формул)[^)]*\)", clean, re.I):
        rep.warn("equations.images", "an image looks like a formula; formulas must be typed in an equation editor "
                                     "(write them as $$...$$)")

    # Citations
    refs = [r for r in (meta.get("references") or []) if str(r).strip()]
    cites = find_citations(clean)
    cited = {n for _, n in cites}
    if refs:
        over = sorted(n for n in cited if n < 1 or n > len(refs))
        if over:
            rep.error("cite.range", f"citation(s) {over} exceed the {len(refs)} entries of the reference list")
        never = [i for i in range(1, len(refs) + 1) if i not in cited]
        if never:
            rep.error("cite.unused", f"reference(s) {never} are never cited in the text")
        seen, expected = [], 1
        for _, n in cites:
            if n in seen:
                continue
            seen.append(n)
            if n != expected:
                rep.error("cite.order", f"[{n}] is cited before [{expected}]; references must be numbered in order of "
                                        "first appearance -- renumber the list and the citations")
                break
            expected += 1
    if re.search(r"\(\s*[A-ZА-ЯІЇЄҐ][\w'-]+,?\s+(?:et al\.,?\s+)?(?:19|20)\d{2}\s*\)", clean) and cites == []:
        rep.error("cite.style", "author-year citations found; use numeric citations in square brackets [1]")

    # Length estimate (a real page count needs --pdf)
    # ~280 words per A4 page at 13 pt / 1.3 spacing with 25 mm margins, plus room for each float/equation.
    top = text_of(meta.get("abstract_top") or meta.get(f"abstract_{lang}"))
    n_words = words(re.sub(r"\$[^$]*\$", " x ", clean)) + words(" ".join(map(str, refs))) + words(top)
    pages = n_words / 280 + 0.4 * len(IMAGE_RE.findall(clean)) + 0.15 * len(table_blocks(body)) \
        + 0.06 * len(DISPLAY_MATH_RE.findall(clean))
    rep.info("length.estimate", f"~{n_words} words incl. title-block abstract and references, roughly {pages:.1f} "
                                "pages; the article must be 5-10 pages (up to 15 by editorial decision) -- "
                                "pass --pdf for an exact count")

    # Language
    if lang == "uk" and re.search(r"(?m)^#\s+(Introduction|Conclusions)\b", clean):
        rep.error("lang.headings", "language is 'uk' but section headings are English")
    if lang == "en" and re.search(r"(?m)^#\s+(Вступ|Висновки)\b", clean):
        rep.error("lang.headings", "language is 'en' but section headings are Ukrainian")


def reminders(meta: dict, rep: Report) -> None:
    rep.info("submit.channel", "submit through the online system at https://asac.kpi.ua (registered account); "
                               "the file must be .doc/.docx -- PDF is not accepted")
    rep.info("submit.attachments", "attach: (1) expert-commission act on the possibility of open publication "
                                   "(акт експертизи, with seal; for authors from Ukraine); (2) permission letters for "
                                   "any borrowed figures/tables; (3) funding and conflict-of-interest statement")
    rep.info("submit.confirm", "the author confirms: unpublished, not under review elsewhere, approved by all co-authors "
                               "and by the institution")
    rep.info("ai.policy", "LLMs cannot be authors; document any generative-AI use in Materials and methods "
                          "(language editing need not be declared)")
    if re.search(r"(?i)human|patient|volunteer|animal|людин|пацієнт|тварин", str(meta)):
        rep.info("ethics", "studies with humans/animals must state compliance with ethical standards at the end of "
                           "Materials and methods")


# ---------------------------------------------------------------- layout checks (docx / pdf)

def check_docx(path: Path, rep: Report) -> None:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    for i, s in enumerate(doc.sections):
        w, h = s.page_width.mm, s.page_height.mm
        if abs(w - PAGE_W_MM) > 1 or abs(h - PAGE_H_MM) > 1:
            rep.error("docx.page", f"section {i + 1}: page is {w:.0f}x{h:.0f} mm, must be A4 (210x297)")
        margins = [s.top_margin.mm, s.bottom_margin.mm, s.left_margin.mm, s.right_margin.mm]
        if any(abs(m - MARGIN_MM) > 0.6 for m in margins):
            rep.error("docx.margins", f"section {i + 1}: margins {[round(m, 1) for m in margins]} mm, must all be 25 mm")
    hy = doc.settings.element.find(qn("w:autoHyphenation"))
    if hy is not None and hy.get(qn("w:val"), "true") not in ("false", "0", "off"):
        rep.error("docx.hyphenation", "automatic hyphenation is on; the journal forbids it")

    def eff(run_el, para, attr):
        rPr = run_el.rPr
        if rPr is not None:
            if attr == "font":
                rf = rPr.find(qn("w:rFonts"))
                if rf is not None and rf.get(qn("w:ascii")):
                    return rf.get(qn("w:ascii"))
            else:
                sz = rPr.find(qn("w:sz"))
                if sz is not None:
                    return int(sz.get(qn("w:val"))) / 2
        st = para.style
        while st is not None:
            v = st.font.name if attr == "font" else (st.font.size.pt if st.font.size else None)
            if v:
                return v
            st = st.base_style
        return None

    bad_font, bad_size, bad_ls, n_runs, n_par = 0, 0, 0, 0, 0
    code_ids = {"SourceCode", "VerbatimChar"}
    for p in doc.paragraphs:
        if p.style is not None and p.style.style_id in code_ids:
            continue
        txt_runs = [r for r in p.runs if r.text.strip()]
        if not txt_runs:
            continue
        n_par += 1
        for r in txt_runs:
            rs = r._r.rPr.find(qn("w:rStyle")) if r._r.rPr is not None else None
            if rs is not None and rs.get(qn("w:val")) in code_ids:
                continue
            n_runs += 1
            if (eff(r._r, p, "font") or "") != FONT:
                bad_font += 1
            size = eff(r._r, p, "size")
            if size is not None and abs(size - FONT_SIZE_PT) > 0.01:
                bad_size += 1
        ls = p.paragraph_format.line_spacing
        if ls is None and p.style is not None:
            st = p.style
            while st is not None and ls is None:
                ls = st.paragraph_format.line_spacing
                st = st.base_style
        if not isinstance(ls, float) or abs(ls - LINE_SPACING) > 0.02:
            bad_ls += 1
    if bad_font:
        rep.error("docx.font", f"{bad_font} of {n_runs} text runs are not in {FONT}")
    if bad_size:
        rep.error("docx.size", f"{bad_size} of {n_runs} text runs are not {FONT_SIZE_PT} pt")
    if bad_ls:
        rep.error("docx.spacing", f"{bad_ls} of {n_par} paragraphs do not use line spacing {LINE_SPACING}")
    if any(p.style is not None and p.style.style_id in code_ids for p in doc.paragraphs):
        rep.warn("docx.code", "the file contains code blocks in a monospace font; keep code listings minimal")
    for t in doc.tables:
        if t._tbl.tblPr.find(qn("w:tblBorders")) is None and not t.style:
            rep.warn("docx.table", "a table has no borders")
    if not any(p.paragraph_format.page_break_before for p in doc.paragraphs):
        rep.warn("docx.endmatter", "no page break found: abstracts/keywords/author info must be on separate pages "
                                   "at the end of the file")


def check_pdf(path: Path, rep: Report) -> None:
    try:
        import pymupdf as fitz
    except ImportError:  # pragma: no cover
        import fitz
    with fitz.open(str(path)) as doc:
        total = doc.page_count
        article = total
        # The end matter starts on a new page with "УДК ..." (the Ukrainian block always comes first);
        # page 0 is skipped because a Ukrainian article also opens with "УДК".
        for i, page in enumerate(doc):
            text = page.get_text()
            if i > 0 and (re.search(r"^\s*УДК\s", text, re.M) or re.search(r"^\s*Бібл\.\s*\d+", text, re.M)):
                article = i
                break
    if article < 5:
        rep.error("pages", f"the article part is {article} page(s); the minimum is 5 ({total} pages incl. end matter)")
    elif article > 15:
        rep.error("pages", f"the article part is {article} pages; the maximum is 10 (15 by editorial decision)")
    elif article > 10:
        rep.warn("pages", f"the article part is {article} pages; over 10 pages needs the editorial board's consent")
    else:
        rep.info("pages", f"the article part is {article} pages (5-10 required); {total} pages incl. end matter")


SOURCE_DIR = Path.cwd()


def main(argv=None) -> int:
    global SOURCE_DIR
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("source", type=Path, help="manuscript Markdown file with YAML front matter")
    ap.add_argument("--docx", type=Path, help="built .docx to check the layout of")
    ap.add_argument("--pdf", type=Path, help="PDF rendering of the .docx (render_preview.py) to count pages")
    ap.add_argument("--year", type=int, default=datetime.date.today().year,
                    help="current year for the 'last 5 years' rule (default: this year)")
    ap.add_argument("--no-reminders", action="store_true", help="omit the submission-checklist INFO lines")
    args = ap.parse_args(argv)
    rep = Report()
    try:
        src = load_source(args.source)
    except (SourceError, OSError) as exc:
        print(f"ERROR source: {exc}")
        print("SUMMARY: 1 error(s), 0 warning(s), 0 info")
        return 1
    SOURCE_DIR = args.source.parent.resolve()
    check_meta(src.meta, rep, args.year)
    check_body(src.meta, src.body, rep)
    if not args.no_reminders:
        reminders(src.meta, rep)
    if args.docx:
        if args.docx.suffix.lower() != ".docx":
            rep.error("docx.format", "the submission must be .doc/.docx")
        elif not args.docx.is_file():
            rep.error("docx.missing", f"not found: {args.docx}")
        else:
            check_docx(args.docx, rep)
    if args.pdf:
        if not args.pdf.is_file():
            rep.error("pdf.missing", f"not found: {args.pdf}")
        else:
            check_pdf(args.pdf, rep)
    rep.print()
    return 1 if rep.count("ERROR") else 0


if __name__ == "__main__":
    sys.exit(main())
