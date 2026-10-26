# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "lxml>=5"]
# ///
"""Behavioral checks for VOTTP formatting and content preservation.

uv run tests/vottp-paper-format/test_vottp.py [--keep /tmp/vottp-test]
The retained fixture is test data, not a journal submission.
"""
from __future__ import annotations

import argparse
import base64
import importlib.util
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

SCRIPT = Path(__file__).resolve().parents[2] / "skills/vottp-paper-format/scripts/vottp_docx.py"
spec = importlib.util.spec_from_file_location("vottp", SCRIPT)
v = importlib.util.module_from_spec(spec)
spec.loader.exec_module(v)


def run(*args, expected=0):
    p = subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], capture_output=True, text=True)
    assert p.returncode == expected, (args, p.returncode, p.stdout, p.stderr)
    return p


def mutate(source, target, callback):
    members, root = v.read_docx(source)
    callback(root)
    with zipfile.ZipFile(target, "w") as z:
        for name, (info, data) in members.items():
            z.writestr(info, v.ET.tostring(root) if name == v.PART else data)


def exercise(work):
    source = work / "draft.docx"
    d = Document()
    d.add_paragraph("УДК 004.42")
    d.add_paragraph("AUTHOR Test")
    d.add_paragraph("Test institution")
    d.add_paragraph("test@example.invalid")
    d.add_paragraph("Formatting regression fixture")
    d.add_paragraph("This fixture exercises formatting and preservation of editable content. " * 30)
    d.add_paragraph("Keywords: formatting, equations")
    d.add_paragraph("Перевірка збереження вмісту документа.")
    d.add_paragraph("Ключові слова: форматування, формули")
    d.add_paragraph("INTRODUCTION")
    p = d.add_paragraph("Text with ")
    p.add_run("bold").bold = True
    p.add_run(" and ")
    p.add_run("italic").italic = True
    p.add_run(" and x")
    p.add_run("2").font.superscript = True
    symbol = p.add_run("α")
    symbol.font.name = "Symbol"
    scientific = d.styles.add_style("Scientific symbols", WD_STYLE_TYPE.CHARACTER)
    scientific.font.name = "Symbol"
    p.add_run("β").style = scientific
    # A hyperlink must survive; flattening p.text would destroy it.
    link = OxmlElement("w:hyperlink")
    rel = d.part.relate_to("https://example.invalid/verified-target", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = " source [1]"
    r.append(t)
    link.append(r)
    p._p.append(link)
    # Nonstandard-but-valid alignment must not crash a high-level enum reader.
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "start")
    p._p.get_or_add_pPr().append(jc)
    eq = d.add_paragraph()
    math = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    mt = OxmlElement("m:t")
    mt.text = "E=mc²"
    mr.append(mt)
    math.append(mr)
    eq._p.append(math)
    d.add_paragraph("Таблиця 1")
    d.add_paragraph("Preserved cell contents")
    table = d.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for cell, value in zip([c for row in table.rows for c in row.cells], ["Name", "Value", "Case A", "42"]):
        cell.text = value
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZlS8AAAAASUVORK5CYII=")
    d.add_picture(io.BytesIO(png), width=Inches(0.2))
    d.add_paragraph("Рис. 1. Small test image")
    d.add_paragraph("References")
    d.add_paragraph("1. Test source. https://example.invalid/verified-target")
    d.sections[0].header.paragraphs[0].text = "Test header retained for preservation testing"
    d.add_section()
    d.add_paragraph("Second section with a numbered list")
    d.add_paragraph("Numbered item", style="List Number")
    d.save(source)
    roles = work / "roles.json"
    run("inspect", source, "--out", roles)
    inventory = json.loads(roles.read_text())
    mapping = {
        0: "udc", 1: "author", 2: "affiliation", 3: "contact", 4: "title",
        5: "abstract-en", 6: "keywords-en", 7: "abstract-uk", 8: "keywords-uk",
        9: "heading", 11: "equation", 12: "table-label", 13: "table-title",
        18: "figure", 19: "figure-caption", 20: "reference-heading", 21: "reference",
    }
    for item in inventory["paragraphs"]:
        item["role"] = mapping.get(item["index"], item["role"])
    v.write_json(roles, inventory)
    before, original_root = v.read_docx(source)
    for profile in ("guide", "published"):
        output = work / f"{profile}.docx"
        report = work / f"{profile}-report.json"
        run("format", source, "--roles", roles, "--out", output, "--profile", profile)
        run("check", output, "--roles", roles, "--original", source, "--out", report, "--profile", profile)
        result = json.loads(report.read_text())
        assert all(result["preservation"].values()) and result["errors"] == 0
        assert result["abstract_characters_including_spaces"]["en"] > 1800
        after, root = v.read_docx(output)
        assert all(before[n][1] == after[n][1] for n in before if n != v.PART)
        assert root.xpath("count(.//w:tbl)", namespaces=v.NS) == 1
        assert root.xpath("count(.//w:hyperlink)", namespaces=v.NS) == 1
        assert root.xpath(".//w:r[w:t='bold']/w:rPr/w:b", namespaces=v.NS)
        assert root.xpath(".//w:r[w:t='α']/w:rPr/w:rFonts/@w:ascii", namespaces=v.NS) == ["Symbol"]
        assert not root.xpath(".//w:r[w:t='β']/w:rPr/w:rFonts", namespaces=v.NS)
        assert v.normalized_content(root) == v.normalized_content(original_root)
        Document(output)  # Reopen with an independent consumer.
    # A malformed role map must never replace an existing output.
    bad = work / "bad-roles.json"
    stale = json.loads(roles.read_text())
    stale["source_sha256"] = "stale"
    v.write_json(bad, stale)
    sentinel = work / "existing.docx"
    sentinel.write_bytes(b"unchanged sentinel")
    run("format", source, "--roles", bad, "--out", sentinel, expected=2)
    assert sentinel.read_bytes() == b"unchanged sentinel"
    stale["source_sha256"] = inventory["source_sha256"]
    stale["paragraphs"][0]["role"] = "invented-role"
    v.write_json(bad, stale)
    run("format", source, "--roles", bad, "--out", sentinel, expected=2)
    assert sentinel.read_bytes() == b"unchanged sentinel"
    run("format", source, "--roles", roles, "--out", source, expected=2)
    # Independent corruptions should be observed by check.
    broken = work / "bad-margin.docx"
    mutate(work / "guide.docx", broken, lambda r: r.xpath(".//w:pgMar", namespaces=v.NS)[0].set(v.q("left"), "100"))
    run("check", broken, "--roles", roles, "--original", source, "--out", work / "bad-margin.json", expected=1)
    def remove_title_bold(r):
        for el in r.xpath(".//w:r[w:t='Formatting regression fixture']/w:rPr/w:b", namespaces=v.NS):
            el.getparent().remove(el)
    mutate(work / "guide.docx", work / "plain-title.docx", remove_title_bold)
    run("check", work / "plain-title.docx", "--roles", roles, "--original", source, "--out", work / "plain-title.json", expected=1)
    def remove_math(r):
        m = r.xpath(".//m:oMath", namespaces=v.NS)[0]
        m.getparent().remove(m)
    mutate(work / "guide.docx", work / "lost-math.docx", remove_math)
    # Changed text/structure invalidates the inventory immediately.
    run("check", work / "lost-math.docx", "--roles", roles, "--original", source, "--out", work / "lost-math.json", expected=2)
    def floating(r):
        r.xpath(".//wp:inline", namespaces=v.NS)[0].tag = f"{{{v.NS['wp']}}}anchor"
    mutate(work / "guide.docx", work / "floating.docx", floating)
    run("check", work / "floating.docx", "--roles", roles, "--original", source, "--out", work / "floating.json", expected=1)
    assert any(f["rule"] == "F05" for f in json.loads((work / "floating.json").read_text())["findings"])
    def tracked(r):
        r.xpath("./w:body/w:p", namespaces=v.NS)[0].append(v.ET.Element(v.q("ins")))
    mutate(source, work / "tracked.docx", tracked)
    run("inspect", work / "tracked.docx", "--out", work / "tracked.json", expected=2)
    # A binary media mutation must fail even if document.xml still matches.
    with zipfile.ZipFile(work / "guide.docx") as z, zipfile.ZipFile(work / "changed-media.docx", "w") as out:
        for info in z.infolist():
            data = z.read(info.filename)
            out.writestr(info, data + b"changed" if info.filename.startswith("word/media/") else data)
    run("check", work / "changed-media.docx", "--roles", roles, "--original", source, "--out", work / "changed-media.json", expected=1)
    print("PASS: both profiles; native math, media, hyperlinks, tables, emphasis, symbols and multiple sections preserved")
    print("PASS: stale/invalid maps, in-place writes, margin defects, missing math, floating images, revisions and media changes detected")


def canonical_document_xml(path):
    """word/document.xml of a .docx in C14N form (independent of zip metadata and serialisation)."""
    import zipfile
    from lxml import etree
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    return etree.tostring(root, method="c14n")


def exercise_examples(work):
    """Shipped worked examples: documented commands run clean and reproduce the shipped outputs (canonical XML)."""
    base = SCRIPT.parents[1] / "assets" / "examples"
    examples = sorted(p for p in base.iterdir() if p.is_dir())
    assert len(examples) >= 2 and {p.name[:3] for p in examples} >= {"en-", "uk-"}, examples
    for ex in examples:
        for name in ("draft.docx", "roles.json", "paper.docx", "layout-report.json", "README.md"):
            assert (ex / name).is_file(), (ex.name, name)
        readme = (ex / "README.md").read_text(encoding="utf-8")
        assert "CC BY 4.0" in readme and "doi.org/10.31891/" in readme, ex.name
        out = work / ex.name
        out.mkdir(parents=True, exist_ok=True)
        run("check", ex / "paper.docx", "--roles", ex / "roles.json", "--original", ex / "draft.docx", "--out", out / "check.json")
        shipped = json.loads((ex / "layout-report.json").read_text())
        assert shipped["errors"] == 0 and all(shipped["preservation"].values()), ex.name
        assert json.loads((out / "check.json").read_text()) == shipped, ex.name
        run("format", ex / "draft.docx", "--roles", ex / "roles.json", "--out", out / "paper.docx")
        # Byte equality is brittle (zip timestamps, lxml serialisation): compare the
        # canonicalised word/document.xml, then require the fresh output to check clean.
        assert canonical_document_xml(out / "paper.docx") == canonical_document_xml(ex / "paper.docx"), ex.name
        run("check", out / "paper.docx", "--roles", ex / "roles.json", "--original", ex / "draft.docx", "--out", out / "check-fresh.json")
        assert json.loads((out / "check-fresh.json").read_text())["errors"] == 0, ex.name
        run("inspect", ex / "draft.docx", "--out", out / "inspect.json")
        fresh = json.loads((out / "inspect.json").read_text())
        roles = json.loads((ex / "roles.json").read_text())
        assert fresh["source_sha256"] == roles["source_sha256"]
        assert [(p["index"], p["text"]) for p in fresh["paragraphs"]] == [(p["index"], p["text"]) for p in roles["paragraphs"]]
        assert all(p["role"] != "body" or not p["in_table"] for p in roles["paragraphs"])
        Document(ex / "paper.docx")
    print("PASS: shipped examples check clean, reproduce the shipped document.xml and keep their role inventories")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--keep", type=Path)
    args = parser.parse_args()
    if args.keep:
        args.keep.mkdir(parents=True, exist_ok=True)
        exercise(args.keep)
        exercise_examples(args.keep)
    else:
        with tempfile.TemporaryDirectory(prefix="vottp-test-") as directory:
            exercise(Path(directory))
            exercise_examples(Path(directory))
