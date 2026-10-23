# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "python-docx>=1.1",
#     "pyyaml>=6",
#     "pypandoc_binary>=1.13",
#     "pymupdf>=1.24",
#     "pillow",
# ]
# ///
"""Plain-assert (no pytest) regression harness for the asac-paper-format skill
scripts: build_paper.py, check_paper.py, import_source.py, render_preview.py.

    uv run tests/asac-paper-format/test_asac.py

Works from any CWD; every path is derived from this file's location or a fresh
temporary directory. The interpreter running this file has every dependency
the skill scripts declare, so it runs them directly (sys.executable) instead of
re-resolving each script's PEP 723 environment.

Covers: the shipped example end-to-end (layout of the produced .docx against
the journal's rules), a Ukrainian-language manuscript, all-or-nothing build
errors, every major check_paper.py rule on deliberately broken sources and a
deliberately broken .docx, zip-level OOXML well-formedness, the importer on the
journal's own sample and on a built file, and (when LibreOffice is installed)
rendering plus the PDF page-count check. Prints a final SUMMARY line and exits
non-zero if anything failed.
"""
from __future__ import annotations

import atexit
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import yaml
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm

REPO_ROOT = Path(__file__).resolve().parents[2]
SKILL_DIR = REPO_ROOT / "skills" / "asac-paper-format"
SCRIPTS = SKILL_DIR / "scripts"
BUILD = SCRIPTS / "build_paper.py"
CHECK = SCRIPTS / "check_paper.py"
IMPORT = SCRIPTS / "import_source.py"
PREVIEW = SCRIPTS / "render_preview.py"
EXAMPLE = SKILL_DIR / "assets" / "example" / "paper.md"
SAMPLE_EN = SKILL_DIR / "assets" / "official" / "sample_eng.docx"
SAMPLE_UA = SKILL_DIR / "assets" / "official" / "sample_ua.docx"
PY = sys.executable
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

_tmp = tempfile.TemporaryDirectory(prefix="asac-paper-format-test-")
atexit.register(_tmp.cleanup)
WORK = Path(_tmp.name)

sys.path.insert(0, str(SCRIPTS))
import asac_common as ac  # noqa: E402
import render_preview as rp  # noqa: E402

RESULTS: list[tuple[str, bool, str]] = []
SKIPPED: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    RESULTS.append((name, bool(cond), detail))
    print(f"[{'PASS' if cond else 'FAIL'}] {name}" + (f" -- {detail}" if detail and not cond else ""))


def skip(name: str, reason: str) -> None:
    SKIPPED.append(name)
    print(f"[SKIP] {name} -- {reason}")


def run(*args, cwd: Path | None = None, timeout: int = 300) -> subprocess.CompletedProcess:
    return subprocess.run([PY, *map(str, args)], cwd=cwd, capture_output=True, text=True, timeout=timeout)


def split_source(path: Path) -> tuple[dict, str]:
    src = ac.load_source(path)
    return src.meta, src.body


def write_source(name: str, meta: dict, body: str, assets_from: Path | None = EXAMPLE.parent) -> Path:
    d = WORK / Path(name).stem
    d.mkdir(parents=True, exist_ok=True)
    if assets_from:
        for f in assets_from.glob("*.png"):
            shutil.copy(f, d / f.name)
    p = d / name
    p.write_text("---\n" + yaml.safe_dump(meta, allow_unicode=True, sort_keys=False) + "---\n" + body,
                 encoding="utf-8")
    return p


def codes(out: str, level: str) -> set[str]:
    return set(re.findall(rf"(?m)^{level}\s+([\w.]+):", out))


def ooxml_problems(path: Path) -> list[str]:
    problems = []
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        for n in names:
            if n.endswith((".xml", ".rels")):
                try:
                    ET.fromstring(z.read(n))
                except ET.ParseError as exc:
                    problems.append(f"{n}: {exc}")
        for n in [x for x in names if x.endswith(".rels")]:
            base = n.replace("_rels/", "").removesuffix(".rels")
            base_dir = base.rsplit("/", 1)[0] if "/" in base else ""
            for rel in ET.fromstring(z.read(n)):
                if rel.get("TargetMode") == "External":
                    continue
                target = rel.get("Target", "")
                full = target.lstrip("/") if target.startswith("/") else \
                    str(Path(base_dir, target)).replace("\\", "/")
                full = re.sub(r"[^/]+/\.\./", "", full)
                if full not in names:
                    problems.append(f"{n}: dangling target {target}")
    return problems


def par_texts(doc) -> list[str]:
    return [p.text for p in doc.paragraphs]


# =====================================================================
print("== 1. asac_common helpers")
check("citation ranges expand", ac.citation_numbers("1–3, 5") == [1, 2, 3, 5])
check("citations found in order", [n for _, n in ac.find_citations("a [2] b [1, 3-4] ![x](y) [5](link)")]
      == [2, 1, 3, 4])
check("reference year ignores access date and DOI",
      ac.reference_year("Web page. – URL: https://x.org/2019 (accessed: 05.05.2025). – 2021.") == 2021)
check("has_doi", ac.has_doi("– DOI: 10.1038/nature14539.") and not ac.has_doi("no identifier 2020"))
a = {"surname_uk": "Петренко", "name_uk": "Іван", "patronymic_uk": "Олегович", "surname_en": "Petrenko",
     "name_en": "Ivan"}
check("author forms", ac.author_short(a, "uk") == "І. Петренко" and ac.author_bib(a, "uk") == "Петренко І.О."
      and ac.author_bib(a, "en") == "Petrenko I." and ac.author_full(a, "uk") == "Петренко Іван Олегович")
check("inline markup", ac.inline_segments("*Xu, R.* Survey / x_y // **B**") ==
      [("Xu, R.", False, True), (" Survey / x_y // ", False, False), ("B", True, False)])

# =====================================================================
print("== 2. build the shipped example (English)")
out_en = WORK / "example.docx"
r = run(BUILD, EXAMPLE, "--out", out_en, "--verbose")
check("example builds", r.returncode == 0 and out_en.is_file(), r.stderr[-800:])
check("verbose log lists numbered objects", all(s in r.stderr for s in ("equation (1)", "equation (2)", "table 1",
                                                                         "figure 1")), r.stderr)
check("example .docx is well-formed OOXML", not ooxml_problems(out_en), str(ooxml_problems(out_en)))
doc = Document(str(out_en))
sec = doc.sections[0]
check("A4 page", abs(sec.page_width.mm - 210) < 0.5 and abs(sec.page_height.mm - 297) < 0.5)
check("25 mm margins", all(abs(m.mm - 25) < 0.2 for m in
                           (sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin)))
hy = doc.settings.element.find(qn("w:autoHyphenation"))
check("auto-hyphenation off", hy is not None and hy.get(qn("w:val")) == "false")
P = doc.paragraphs
check("title block: UDC", P[0].text == "UDC 004.852" and P[0].alignment == 0)
check("title block: authors bold right", P[1].text == "I. Petrenko, O. Kovalenko" and P[1].alignment == 2
      and all(r.bold for r in P[1].runs))
check("title block: title upper/bold/centred", P[2].text.isupper() and P[2].alignment == 1
      and all(r.bold for r in P[2].runs))
check("title block: abstract label italic", P[3].runs[0].text.startswith("Abstract") and P[3].runs[0].italic)
kw = next(p for p in P if p.text.startswith("Keywords:"))
check("title block: keywords", kw.runs[0].italic and kw.text.endswith("process monitoring."))
texts = par_texts(doc)
body_runs = [r for p in P for r in p.runs if r.text.strip()]
check("every text run is Times New Roman 13",
      all(r.font.name == "Times New Roman" and r.font.size and r.font.size.pt == 13 for r in body_runs),
      str({(r.font.name, r.font.size.pt if r.font.size else None) for r in body_runs}))
ls = {round(p.paragraph_format.line_spacing, 2) for p in P if p.text.strip()
      and p.paragraph_format.line_spacing is not None}
check("line spacing 1.3 everywhere outside tables", ls == {1.3}, str(ls))
eqs = [p for p in P if p._p.find(f"{{{M_NS}}}oMath") is not None and re.search(r"\(\d+\)$", p.text)]
check("two numbered native equations", [p.text.strip()[-3:] for p in eqs] == ["(1)", "(2)"],
      str([p.text for p in eqs]))
check("equations use centre+right tabs", all(len(p._p.findall(f".//{qn('w:tab')}")) >= 2 for p in eqs))
where = [p for p in P if p.text.startswith("where")]
check("'where' lines flush left", where and all(p.paragraph_format.first_line_indent == 0 for p in where))
fig = next((p for p in P if p.text.startswith("Figure 1.")), None)
check("figure caption 'Figure 1.' italic label, centred", fig is not None and fig.runs[0].italic
      and fig.alignment == 1)
tab_i = next((i for i, p in enumerate(P) if p.text == "Table 1."), None)
check("table label right-aligned italic + centred title", tab_i is not None and P[tab_i].alignment == 2
      and P[tab_i].runs[0].italic and P[tab_i + 1].alignment == 1)
tbl = doc.tables[0]
check("table fully ruled", tbl._tbl.tblPr.find(qn("w:tblBorders")) is not None)
check("table text 13 pt TNR", all(r.font.size.pt == 13 and r.font.name == "Times New Roman"
                                  for c in tbl._cells for p in c.paragraphs for r in p.runs if r.text.strip()))
ref_head = texts.index("REFERENCES")
refs = [t for t in texts[ref_head + 1:] if re.match(r"^\d+\. ", t)]
check("12 numbered references", len(refs) == 12 and refs[0].startswith("1. LeCun"), str(refs[:2]))
check("heading author italic in references", P[ref_head + 1].runs[1].italic)
check("disclosure section before references", texts.index("Disclosure of interests") < ref_head)
breaks = [p.text for p in P if p.paragraph_format.page_break_before]
check("end matter on three new pages", breaks[:2] == ["УДК 004.852", "UDC 004.852"]
      and "Про авторів / About the authors" in breaks, str(breaks))
check("end matter counts computed", "Бібл. 12, іл. 1, табл. 1" in texts and "Ref. 12, fig. 1, tab. 1" in texts)
check("UA citation line", any(t.startswith("Адаптивний градієнтний") and "Петренко І.О." in t
                              and "– 2026. – № 2 (51). – с." in t for t in texts))
check("structured UA abstract labels", any(t.startswith("Анотація: Проблематика.") for t in texts)
      and any(t.startswith("Методика реалізації.") for t in texts))
check("author info: ORCID, corresponding", "ORCID: 0000-0000-0000-0000" in texts
      and "Відповідальний автор / Corresponding author" in texts)
check("no header/footer text added", all(not (s.header.paragraphs and s.header.paragraphs[0].text.strip())
                                         for s in doc.sections))
with zipfile.ZipFile(out_en) as z:
    numbering = z.read("word/numbering.xml").decode()
check("bullets are en dashes", 'w:lvlText w:val="–"' in numbering)

# =====================================================================
print("== 3. check the shipped example")
r = run(CHECK, EXAMPLE, "--docx", out_en, "--year", "2026")
check("example has no ERROR", r.returncode == 0 and not codes(r.stdout, "ERROR"), r.stdout[-1500:])
check("summary line printed", "SUMMARY: 0 error(s)" in r.stdout)
check("placeholders reported", "authors.placeholder" in codes(r.stdout, "WARN"))
check("missing DOI reported", "refs.doi" in codes(r.stdout, "WARN"))
check("submission reminders", {"submit.channel", "submit.attachments", "ai.policy"} <= codes(r.stdout, "INFO"))

# =====================================================================
print("== 4. Ukrainian manuscript")
meta, body = split_source(EXAMPLE)
meta_uk = dict(meta, language="uk", references=meta["references"][:3])
body_uk = """
# Вступ

Табличні дані [1] є поширеними в задачах керування [2]. Метою роботи є побудова моделі [3].

# Матеріали і методи

Модель задається формулою

$$ y = a x + b $$

де $a$ і $b$ -- коефіцієнти. Результати подано на рис. 1 і в табл. 1.

![Похибка моделі на валідаційній вибірці](fig1.png){width=10cm}

Table: Параметри моделі

| Параметр | Значення |
|---|---|
| a | 1 |

# Висновки

Запропоновано модель.
"""
src_uk = write_source("paper_uk.md", meta_uk, body_uk)
out_uk = WORK / "paper_uk.docx"
r = run(BUILD, src_uk, "--out", out_uk)
check("Ukrainian manuscript builds", r.returncode == 0, r.stderr[-600:])
if out_uk.is_file():
    d = Document(str(out_uk))
    t = par_texts(d)
    check("УДК / Анотація / Ключові слова in title block", t[0].startswith("УДК") and
          t[3].startswith("Анотація:") and any(x.startswith("Ключові слова:") for x in t[:10]))
    check("Рисунок 1. / Таблиця 1.", any(x.startswith("Рисунок 1.") for x in t) and "Таблиця 1." in t)
    check("СПИСОК ВИКОРИСТАНИХ ДЖЕРЕЛ", "СПИСОК ВИКОРИСТАНИХ ДЖЕРЕЛ" in t)
    de = [p for p in d.paragraphs if p.text.startswith("де ")]
    check("'де' line (separate paragraph) flush left", de and de[0].paragraph_format.first_line_indent == 0)
    deyaki = [p for p in d.paragraphs if p.text.startswith("Модель задається")]
    # Word stores indents in twips (709 = 12.507 mm), so compare in mm with a tolerance.
    check("ordinary paragraph keeps its indent",
          deyaki and abs(deyaki[0].paragraph_format.first_line_indent.mm - 12.5) < 0.1)
r = run(CHECK, src_uk, "--year", "2026")
check("UA check flags short conclusions", "conclusions.length" in codes(r.stdout, "WARN"), r.stdout[-800:])

# =====================================================================
print("== 5. build errors are all-or-nothing")


def build_fails(name: str, meta_: dict, body_: str, needle: str, out_name: str | None = None) -> None:
    src = write_source(name, meta_, body_)
    out = WORK / (out_name or (Path(name).stem + ".docx"))
    rr = run(BUILD, src, "--out", out)
    check(f"build error: {needle}", rr.returncode == 1 and needle in rr.stderr and not out.exists(),
          f"rc={rr.returncode} stderr={rr.stderr[-400:]}")


m = dict(meta)
m.pop("abstract_uk")
build_fails("no_abs.md", m, body, "abstract_uk")
build_fails("bad_lang.md", dict(meta, language="de"), body, "language")
build_fails("no_authors.md", dict(meta, authors=[]), body, "authors")
build_fails("no_refs.md", dict(meta, references=[]), body, "references")
build_fails("missing_img.md", meta, body.replace("fig1.png", "nope.png"), "pandoc")
nofm = WORK / "nofm.md"
nofm.write_text("# Introduction\n\ntext\n", encoding="utf-8")
r = run(BUILD, nofm, "--out", WORK / "nofm.docx")
check("build error: no front matter", r.returncode == 1 and "front matter" in r.stderr)
r = run(BUILD, EXAMPLE, "--out", WORK / "x.pdf")
check("build refuses PDF output", r.returncode == 1 and "PDF" in r.stderr and not (WORK / "x.pdf").exists())

# =====================================================================
print("== 6. checker catches broken sources")


def check_codes(name: str, meta_: dict, body_: str) -> str:
    src = write_source(name, meta_, body_)
    return run(CHECK, src, "--year", "2026", "--no-reminders").stdout


cases = [
    ("cite_range.md", meta, body.replace("[10]", "[13]"), "ERROR", "cite.range"),
    ("cite_order.md", meta, body.replace("tabular sensor records [1]", "tabular sensor records [2]")
     .replace("decision\ntrees [2]", "decision\ntrees [1]"), "ERROR", "cite.order"),
    ("intro_float.md", meta, body.replace("# Materials and methods", "![Early figure](fig1.png)\n\n"
                                                                     "# Materials and methods"), "ERROR",
     "intro.floats"),
    ("no_orcid.md", dict(meta, authors=[{k: v for k, v in meta["authors"][0].items() if k != "orcid"}]), body,
     "ERROR", "authors.orcid"),
    ("few_kw.md", dict(meta, keywords_en=["a b", "c d", "e f", "g h"]), body, "ERROR", "keywords.en.count"),
    ("long_abs.md", dict(meta, abstract_en=dict(meta["abstract_en"], results=meta["abstract_en"]["results"] * 6)),
     body, "ERROR", "abstract.en.chars"),
    ("plain_abs.md", dict(meta, abstract_en="Just one paragraph."), body, "WARN", "abstract.en.structure"),
    ("unpub.md", dict(meta, references=meta["references"] + ["Ivanenko I. Personal communication, 2025."]),
     body + "\nSee also [13].\n", "ERROR", "refs.unpublished"),
    ("todo.md", dict(meta, udc="TODO"), body, "ERROR", "meta.todo"),
    ("no_intro.md", meta, body.replace("# Introduction", "# Background"), "ERROR", "body.introduction"),
    ("no_caption.md", meta, body.replace("![Validation error of the reweighted ensemble over boosting rounds]",
                                         "![]"), "ERROR", "figure.caption"),
    ("no_tab_caption.md", meta, body.replace("Table: Hyper-parameters shared by all compared models\n", ""),
     "ERROR", "table.caption"),
    ("body_refs.md", meta, body + "\n# References\n\n1. x\n", "ERROR", "body.generated"),
    ("no_corr.md", dict(meta, authors=[dict(a, corresponding=False) for a in meta["authors"]]), body, "ERROR",
     "authors.corresponding"),
    ("funding.md", dict(meta, funding="Funded by X."), body, "WARN", "funding.prefix"),
    ("review.md", dict(meta, article_type="review", funding=None), body, "ERROR", "review.disclosure"),
    ("title_abbr.md", dict(meta, title_en="An XGBDT-based method"), body, "WARN", "title.en.abbreviations"),
]
for name, m_, b_, level, code in cases:
    out = check_codes(name, m_, b_)
    check(f"{level} {code}", code in codes(out, level), out[-600:])

# =====================================================================
print("== 7. checker catches a broken .docx")
broken = WORK / "broken.docx"
d = Document(str(out_en))
d.sections[0].left_margin = Mm(20)
d.settings.element.find(qn("w:autoHyphenation")).set(qn("w:val"), "true")
d.paragraphs[10].runs[0].font.name = "Arial"
d.paragraphs[11].paragraph_format.line_spacing = 1.0
d.save(str(broken))
r = run(CHECK, EXAMPLE, "--docx", broken, "--no-reminders")
got = codes(r.stdout, "ERROR")
check("docx errors detected", {"docx.margins", "docx.hyphenation", "docx.font", "docx.spacing"} <= got, str(got))
r = run(CHECK, EXAMPLE, "--docx", WORK / "nothing.pdf", "--no-reminders")
check("PDF submission rejected", "docx.format" in codes(r.stdout, "ERROR"))

# =====================================================================
print("== 8. importer")
imp = WORK / "imp" / "paper.md"
r = run(IMPORT, SAMPLE_EN, "--out", imp)
check("imports the journal's English sample", r.returncode == 0 and imp.is_file(), r.stderr[-500:])
if imp.is_file():
    im, ib = split_source(imp)
    check("sample: UDC and 4 references picked up", im["udc"] == "004.05" and len(im["references"]) == 4,
          f"{im.get('udc')} {len(im.get('references', []))}")
    check("sample: citations unescaped", "[1]" in ib and "\\[" not in ib)
    check("sample: image extracted", (imp.parent / "media" / "image1.png").is_file())
    check("sample: TODOs left for the checker", "meta.todo" in codes(run(CHECK, imp).stdout, "ERROR"))
r = run(IMPORT, SAMPLE_EN, "--out", imp)
check("importer refuses to overwrite", r.returncode == 1 and "--force" in r.stderr)
imp2 = WORK / "imp2" / "paper.md"
r = run(IMPORT, out_en, "--out", imp2)
if imp2.is_file():
    im2, ib2 = split_source(imp2)
    check("round trip: 12 references", len(im2["references"]) == 12, str(len(im2["references"])))
    check("round trip: numbered equations become display math", len(re.findall(r"(?m)^\$\$", ib2)) == 2)
    check("round trip: keywords", len(im2["keywords_en"]) == 6)
else:
    check("round trip import", False, r.stderr[-500:])
r = run(IMPORT, SAMPLE_UA, "--out", WORK / "imp3" / "paper.md")
check("imports the Ukrainian sample as uk", r.returncode == 0 and "language=uk" in r.stderr, r.stderr[-300:])

# =====================================================================
print("== 9. render + page count")
try:
    rp.find_soffice()
    have_soffice = True
except RuntimeError:
    have_soffice = False
if not have_soffice:
    skip("render_preview / page count", "LibreOffice (soffice) not found")
else:
    prev = WORK / "preview"
    r = run(PREVIEW, out_en, "--out", prev, timeout=400)
    pdf = prev / "example.pdf"
    check("render produces PDF, pages and contact sheet", r.returncode == 0 and pdf.is_file()
          and (prev / "contact-sheet.png").is_file() and (prev / "page-01.png").is_file(), r.stderr[-500:])
    if pdf.is_file():
        r = run(CHECK, EXAMPLE, "--docx", out_en, "--pdf", pdf, "--no-reminders")
        m_ = re.search(r"article part is (\d+) pages?", r.stdout)
        check("page count reported for the article part", m_ is not None, r.stdout[-400:])
        check("example article within 5-10 pages", m_ is not None and 5 <= int(m_.group(1)) <= 10
              and "pages" not in codes(r.stdout, "ERROR"), m_.group(0) if m_ else "")

# =====================================================================
n_fail = sum(1 for _, ok, _ in RESULTS if not ok)
n_pass = len(RESULTS) - n_fail
print(f"SUMMARY: {n_pass}/{len(RESULTS)} passed, {n_fail} failed, {len(SKIPPED)} skipped")
sys.exit(1 if n_fail else 0)
