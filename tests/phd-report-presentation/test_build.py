# /// script
# requires-python = ">=3.10,<3.14"
# dependencies = [
#     "python-pptx>=1.0",
#     "pillow",
#     "pymupdf>=1.24",
#     "lxml>=5.0",
# ]
# ///
"""Plain-assert (no pytest) regression harness for the phd-report-presentation
skill scripts: build_deck.py, render_preview.py, make_template.py.

Run it with uv (resolves its own dependencies from the PEP 723 header above):

    uv run tests/phd-report-presentation/test_build.py

...or with any interpreter that already has python-pptx/pillow/pymupdf/lxml
installed:

    /path/to/python tests/phd-report-presentation/test_build.py

Works from any CWD -- every path below is derived from this file's own
location (REPO_ROOT/SKILL_DIR) or from a fresh tempfile.TemporaryDirectory(),
never from a hardcoded absolute path.

Builds the example spec and a battery of adversarial specs, opens outputs
with python-pptx, and does zip-level OOXML sanity checks (well-formed XML,
Content_Types coverage, no dangling relationship targets). Prints a final
PASS/FAIL/SKIP summary and exits non-zero if anything failed.

soffice (LibreOffice) is optional: render_preview.py checks need it to
rasterize a deck, and are skipped cleanly (printed as [SKIP]) when `soffice`
cannot be found.
"""
from __future__ import annotations

import atexit
import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

REPO_ROOT = Path(__file__).resolve().parents[2]
SKILL_DIR = REPO_ROOT / "skills" / "phd-report-presentation"
BUILD = SKILL_DIR / "scripts" / "build_deck.py"
PREVIEW = SKILL_DIR / "scripts" / "render_preview.py"
MAKE_TEMPLATE = SKILL_DIR / "scripts" / "make_template.py"
EXAMPLE_SPEC = SKILL_DIR / "assets" / "example-spec.json"
TEMPLATE = SKILL_DIR / "assets" / "template.pptx"

# The interpreter currently running *this* test file. Launched via `uv run`
# it has python-pptx/pillow/pymupdf/lxml installed per the PEP 723 header
# above; launched with some other venv python it's expected to already have
# them (see the module docstring). Either way that's a superset of what each
# individual (also PEP 723) skill script needs, so it doubles as the
# interpreter used to invoke build_deck.py/render_preview.py/make_template.py
# directly below. `uv` itself (found on PATH) is used separately in section 2
# to exercise the scripts' own PEP 723 `uv run` entry point.
PY = sys.executable
UV = shutil.which("uv")

# A fresh, process-local scratch dir for every build output / work file this
# test produces -- replaces the old hardcoded scratchpad path. Cleaned up at
# interpreter exit (also covers the sys.exit() calls at the bottom).
_tmpdir = tempfile.TemporaryDirectory(prefix="phd-report-presentation-test-")
atexit.register(_tmpdir.cleanup)
WORK = Path(_tmpdir.name) / "work"
WORK.mkdir(parents=True)

sys.path.insert(0, str(SKILL_DIR / "scripts"))
import build_deck as bd  # noqa: E402
import render_preview as rp  # noqa: E402


def _soffice_available() -> bool:
    try:
        rp.find_soffice()
        return True
    except RuntimeError:
        return False


HAVE_SOFFICE = _soffice_available()

RESULTS: list[tuple[str, bool, str]] = []
SKIPPED: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    RESULTS.append((name, bool(cond), detail))
    status = "PASS" if cond else "FAIL"
    print(f"[{status}] {name}" + (f" -- {detail}" if detail and not cond else ""))


def skip(name: str, reason: str) -> None:
    SKIPPED.append(name)
    print(f"[SKIP] {name} -- {reason}")


def run(cmd: list[str], cwd: Path | None = None, timeout: int = 180) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, cwd=cwd, capture_output=True, text=True, timeout=timeout)


def write_spec(name: str, obj: dict) -> Path:
    p = WORK / name
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
    return p


# =====================================================================
# Zip-level OOXML sanity checker
# =====================================================================
XML_EXTS = {".xml", ".rels"}
NS_CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"
NS_R = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def check_ooxml_zip(path: Path) -> list[str]:
    """Returns a list of problems (empty == clean)."""
    problems = []
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        # 1. every XML-ish part is well-formed.
        for n in names:
            if Path(n).suffix.lower() in XML_EXTS or n == "[Content_Types].xml":
                try:
                    ET.fromstring(z.read(n))
                except ET.ParseError as exc:
                    problems.append(f"malformed XML in {n}: {exc}")

        # 2. [Content_Types].xml covers every part.
        ct_root = ET.fromstring(z.read("[Content_Types].xml"))
        defaults = {el.get("Extension", "").lower() for el in ct_root.findall(f"{NS_CT}Default")}
        overrides = {el.get("PartName", "").lstrip("/") for el in ct_root.findall(f"{NS_CT}Override")}
        for n in names:
            if n == "[Content_Types].xml":
                continue
            # Path(...).suffix on a dotfile like "_rels/.rels" is '' (pathlib
            # treats ".rels" as an extension-less hidden filename) -- handle
            # the ".rels"/".xml" OPC convention explicitly instead of relying
            # on Path.suffix for the whole name.
            ext = n.rsplit(".", 1)[-1].lower() if "." in Path(n).name else ""
            if n in overrides or ext in defaults:
                continue
            problems.append(f"part {n} not covered by [Content_Types].xml (no Default/Override)")

        # 3. no relationship points at a missing part.
        for n in names:
            if not n.endswith(".rels"):
                continue
            rel_root = ET.fromstring(z.read(n))
            if "/_rels/" in n:
                base_dir, fname = n.split("/_rels/")
                base_part = f"{base_dir}/{fname[:-5]}" if base_dir else fname[:-5]
            else:
                base_part = ""
            for rel in rel_root.findall(f"{NS_R}Relationship"):
                if rel.get("TargetMode") == "External":
                    continue
                target = rel.get("Target", "")
                if target.startswith("http"):
                    continue
                if target.startswith("/"):
                    resolved = target.lstrip("/")
                else:
                    base = base_part.rsplit("/", 1)[0] if "/" in base_part else ""
                    combined = f"{base}/{target}" if base else target
                    out = []
                    for seg in combined.split("/"):
                        if seg == "..":
                            if out:
                                out.pop()
                        elif seg in (".", ""):
                            continue
                        else:
                            out.append(seg)
                    resolved = "/".join(out)
                if resolved not in names:
                    problems.append(f"{n}: relationship {rel.get('Id')} -> missing part '{resolved}'")
    return problems


def shape_ids_per_slide(pptx_path: Path):
    """Returns {slide_index(1-based): [cNvPr ids]} via raw XML (python-pptx
    independent, so this also validates the *saved* XML, not just the
    in-memory object model)."""
    out = {}
    with zipfile.ZipFile(pptx_path) as z:
        slide_names = sorted(
            [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")],
            key=lambda n: int("".join(c for c in Path(n).stem if c.isdigit())),
        )
        for i, n in enumerate(slide_names, start=1):
            root = ET.fromstring(z.read(n))
            ids = [el.get("id") for el in root.iter() if el.tag.endswith("}cNvPr")]
            out[i] = ids
    return out


def count_slidenum_fields(pptx_path: Path) -> dict:
    out = {}
    with zipfile.ZipFile(pptx_path) as z:
        slide_names = sorted(
            [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")],
            key=lambda n: int("".join(c for c in Path(n).stem if c.isdigit())),
        )
        for i, n in enumerate(slide_names, start=1):
            data = z.read(n).decode("utf-8")
            out[i] = data.count('type="slidenum"')
    return out


def empty_placeholder_count(pptx_path: Path) -> int:
    """Count placeholder shapes whose text frame is completely empty (all
    runs empty/absent) -- these are the 'auto-cloned, never filled in'
    placeholders the generator is supposed to always drop."""
    from pptx import Presentation
    prs = Presentation(str(pptx_path))
    n = 0
    for slide in prs.slides:
        for shape in slide.placeholders:
            if not shape.has_text_frame:
                continue
            text = "".join(r.text for p in shape.text_frame.paragraphs for r in p.runs)
            if text.strip() == "":
                n += 1
    return n


# =====================================================================
# 1. Baseline: example spec builds cleanly, deck is structurally sound
# =====================================================================
print("\n=== 1. Baseline example-spec build ===")
out1 = WORK / "example.pptx"
p = run([PY, str(BUILD), "--spec", str(EXAMPLE_SPEC), "--out", str(out1), "--verbose"])
check("example spec builds (exit 0)", p.returncode == 0, p.stderr)
check("example spec: no WARN on clean input", "WARN" not in p.stderr, p.stderr)
check("example.pptx written", out1.is_file())

from pptx import Presentation  # noqa: E402

prs1 = Presentation(str(out1))
with open(EXAMPLE_SPEC, encoding="utf-8") as f:
    example_spec_obj = json.load(f)
n_expected = len(example_spec_obj["slides"])
check("slide count == %d" % n_expected, len(list(prs1.slides)) == n_expected,
      f"got {len(list(prs1.slides))}")

ooxml_problems = check_ooxml_zip(out1)
check("zip-level OOXML sanity (well-formed, Content_Types, no dangling rels)",
      not ooxml_problems, "; ".join(ooxml_problems[:5]))

ids_per_slide = shape_ids_per_slide(out1)
dup_slides = [i for i, ids in ids_per_slide.items() if len(ids) != len(set(ids))]
check("shape ids unique per slide (//p:cNvPr/@id)", not dup_slides, f"dup on slides {dup_slides}")

slidenum_counts = count_slidenum_fields(out1)
# Per spec: slide 1 (title) and the last slide (back_cover) must NOT have a
# slide-number field; every other ("content") slide must have EXACTLY one.
bad_slidenum = {
    i: c for i, c in slidenum_counts.items()
    if (i in (1, n_expected) and c != 0) or (i not in (1, n_expected) and c != 1)
}
check("exactly one slidenum field on content slides, none on title/back_cover",
      not bad_slidenum, f"{bad_slidenum}")

empties = empty_placeholder_count(out1)
check("no empty placeholders left on any slide", empties == 0, f"{empties} empty placeholder(s)")

n_pics = sum(1 for slide in prs1.slides for shape in slide.shapes if shape.shape_type == 13)
check("images present in output (title logos + table/column/image pics)", n_pics > 0, f"{n_pics} pics")


# =====================================================================
# 2. uv run works, from a directory outside the repo, for all 3 scripts
# =====================================================================
print("\n=== 2. uv run (PEP 723) from outside the repo ===")
if UV is None:
    skip("uv run build_deck.py from outside repo", "uv not found on PATH")
    skip("uv run render_preview.py from outside repo", "uv not found on PATH")
    skip("contact sheet produced via uv run", "uv not found on PATH")
    skip("uv run make_template.py from outside repo", "uv not found on PATH")
else:
    outside_dir = Path(tempfile.mkdtemp(prefix="uv-outside-"))
    out_uv = outside_dir / "deck-uv.pptx"
    p = run([UV, "run", str(BUILD), "--spec", str(EXAMPLE_SPEC), "--out", str(out_uv)], cwd=outside_dir)
    check("uv run build_deck.py from outside repo", p.returncode == 0 and out_uv.is_file(), p.stderr[-800:])

    if HAVE_SOFFICE:
        p = run([UV, "run", str(PREVIEW), str(out_uv), "--out", str(outside_dir / "prev")], cwd=outside_dir)
        check("uv run render_preview.py from outside repo", p.returncode == 0, p.stderr[-800:])
        check("contact sheet produced via uv run", (outside_dir / "prev" / "contact-sheet.png").is_file())
    else:
        skip("uv run render_preview.py from outside repo", "soffice not on PATH")
        skip("contact sheet produced via uv run", "soffice not on PATH")

    mt_out = outside_dir / "template-uv.pptx"
    p = run([UV, "run", str(MAKE_TEMPLATE), "--source", str(TEMPLATE), "--out", str(mt_out)], cwd=outside_dir)
    check("uv run make_template.py from outside repo", p.returncode == 0 and mt_out.is_file(), p.stderr[-800:])
    shutil.rmtree(outside_dir, ignore_errors=True)


# =====================================================================
# 3. Adversarial specs -- each should fail cleanly (SpecError -> exit 1,
#    one-line stderr "error: ...", NOT a raw traceback) or WARN+succeed.
# =====================================================================
print("\n=== 3. Adversarial specs ===")


def base_spec(slides):
    return {"meta": {"dissertation_title": "T", "department": "D",
                      "student_name": "S", "supervisor_name": "P"}, "slides": slides}


def expect_clean_error(name: str, spec_obj: dict, must_contain: str = "", cwd: Path | None = None,
                        fname: str | None = None):
    sp = write_spec(fname or (name.replace(" ", "_") + ".json"), spec_obj)
    out = WORK / (name.replace(" ", "_") + ".pptx")
    p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out)], cwd=cwd)
    clean = (p.returncode == 1 and p.stderr.strip().startswith("error:")
             and "Traceback" not in p.stderr)
    ok = clean and (must_contain.lower() in p.stderr.lower() if must_contain else True)
    check(f"adversarial: {name}", ok, f"exit={p.returncode} stderr={p.stderr.strip()[:300]!r}")


def expect_success(name: str, spec_obj: dict, must_warn: str = "", cwd: Path | None = None):
    sp = write_spec(name.replace(" ", "_") + ".json", spec_obj)
    out = WORK / (name.replace(" ", "_") + ".pptx")
    p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out)], cwd=cwd)
    ok = p.returncode == 0 and out.is_file()
    if ok and must_warn:
        ok = must_warn.lower() in p.stderr.lower()
    check(f"adversarial (should succeed): {name}", ok,
          f"exit={p.returncode} stderr={p.stderr.strip()[:400]!r}")
    return p, out


# -- unknown slide type
expect_clean_error("unknown slide type", base_spec([{"type": "nope"}]), "unknown slide type")

# -- empty slides list
expect_clean_error("empty slides list", {"meta": {}, "slides": []}, "non-empty")

# -- missing required field (bullets without items)
expect_clean_error("missing items field", base_spec([{"type": "bullets", "title": "T"}]),
                    "missing required field")

# -- ragged table rows (fewer cells)
expect_clean_error(
    "ragged table fewer cells",
    base_spec([{"type": "table", "title": "T", "columns": ["A", "B"], "rows": [["x"]]}]),
    "cells, expected",
)
# -- ragged table rows (more cells)
expect_clean_error(
    "ragged table more cells",
    base_spec([{"type": "table", "title": "T", "columns": ["A"], "rows": [["x", "y"]]}]),
    "cells, expected",
)
# -- table row not a list
expect_clean_error(
    "table row not a list",
    base_spec([{"type": "table", "title": "T", "columns": ["A"], "rows": ["not-a-row"]}]),
    "must be an array of cells",
)
# -- table cell is a dict (not coercible)
expect_clean_error(
    "table cell is object",
    base_spec([{"type": "table", "title": "T", "columns": ["A"], "rows": [[{"x": 1}]]}]),
    "cell must be",
)
# -- table 0 rows
expect_clean_error(
    "table zero rows",
    base_spec([{"type": "table", "title": "T", "columns": ["A"], "rows": []}]),
    "non-empty array",
)

# -- missing image file
expect_clean_error(
    "missing image file",
    base_spec([{"type": "image", "title": "T", "image": "does-not-exist.png"}]),
    "image not found",
)
# -- unsupported image extension
bad_img = WORK / "chart.svg"
bad_img.write_text("<svg></svg>", encoding="utf-8")
expect_clean_error(
    "unsupported image extension",
    base_spec([{"type": "image", "title": "T", "image": "chart.svg"}]),
    "unsupported image extension",
)

# -- nested sub depth > 2
expect_clean_error(
    "sub nesting depth 3",
    base_spec([{"type": "bullets", "title": "T", "items": [
        {"text": "H", "sub": [{"text": "S1", "sub": ["too deep"]}]}
    ]}]),
    "nesting deeper than 2 levels",
)

# -- non-string bullet item
expect_clean_error(
    "non-string bullet item",
    base_spec([{"type": "bullets", "title": "T", "items": [42]}]),
    "must be a string",
)
# -- boolean bullet item
expect_clean_error(
    "boolean bullet item",
    base_spec([{"type": "bullets", "title": "T", "items": [True]}]),
    "must be a string",
)
# -- non-string title (number)
expect_clean_error(
    "numeric title field",
    base_spec([{"type": "bullets", "title": 123, "items": ["x"]}]),
    "must be a string",
)
# -- bad style enum (typo)
expect_clean_error(
    "style typo",
    base_spec([{"type": "bullets", "title": "T", "items": ["x"], "style": "numbererd"}]),
    "must be one of",
)
# -- two_column wrong column count
expect_clean_error(
    "two_column 3 columns",
    base_spec([{"type": "two_column", "title": "T", "columns": [
        {"sections": [{"items": ["a"]}]}, {"sections": [{"items": ["a"]}]}, {"sections": [{"items": ["a"]}]}
    ]}]),
    "exactly 2 items",
)

# -- 0-row table is already covered; ragged covered. Now "empty string" required field:
expect_clean_error(
    "empty-string title",
    base_spec([{"type": "bullets", "title": "", "items": ["x"]}]),
    "missing required field",
)

print("\n=== 3b. Specs that should succeed with a WARN (never hard-fail) ===")

# -- unknown key / typo detection
expect_success(
    "unknown key typo",
    base_spec([{"type": "bullets", "tilte": "oops", "title": "T", "items": ["x"]}]),
    must_warn="unrecognised field",
)

# -- title slide missing meta fields entirely -> WARN, not crash
expect_success(
    "title missing meta",
    {"meta": {}, "slides": [{"type": "title"}, {"type": "thanks"}, {"type": "back_cover"}]},
    must_warn="missing meta/slide field",
)

# -- back_cover / thanks with extra (unknown) fields -> WARN, not fail
expect_success(
    "back_cover extra fields",
    base_spec([{"type": "thanks"}, {"type": "back_cover", "subtitle": "nope"}]),
    must_warn="unrecognised field",
)

# -- None values in optional fields (table caption null, thanks text null)
expect_success(
    "none values in optional fields",
    base_spec([
        {"type": "table", "title": "T", "caption": None, "columns": ["A"], "rows": [[None]]},
        {"type": "thanks", "text": None},
    ]),
)

# -- numbers/booleans in table cells -> coerced
p_num, out_num = expect_success(
    "numbers booleans in table cells",
    base_spec([{"type": "table", "title": "T", "columns": ["A", "B"], "rows": [[42, True]]}]),
)
if out_num.is_file():
    prs_num = Presentation(str(out_num))
    cell_texts = []
    for slide in prs_num.slides:
        for shape in slide.shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        cell_texts.append(cell.text)
    check("numeric/boolean cells coerced to visible text ('42'/'True'), not crashed",
          "42" in cell_texts and "True" in cell_texts, f"{cell_texts}")

# -- table with >11 rows must not shrink row height below the documented
#    0.75cm floor, and MUST warn that it overflows.
many_rows = [[f"r{i}", "x", "y"] for i in range(16)]
p_many, out_many = expect_success(
    "table many rows floor",
    base_spec([{"type": "table", "title": "T", "columns": ["A", "B", "C"], "rows": many_rows}]),
    must_warn="tight against the available",
)
if out_many.is_file():
    prs_many = Presentation(str(out_many))
    row_h_cm = None
    for slide in prs_many.slides:
        for shape in slide.shapes:
            if shape.has_table:
                row_h_cm = shape.table.rows[0].height / 360000.0
    check("16-row table keeps row height >= 0.75cm floor (no silent shrink)",
          row_h_cm is not None and row_h_cm >= 0.75 - 1e-6, f"row_h_cm={row_h_cm}")

# -- Unicode / Cyrillic / apostrophe variants / guillemets / long word
expect_success(
    "unicode and edge chars",
    base_spec([
        {"type": "bullets", "title": "Юнікод: Кіт у мішку «тест» ʼ'’", "items": [
            "Звичайний пункт з апострофом: Слов'янська ніч",
            "A" * 300,  # one very long "word" (no spaces) -- must not crash wrapping heuristic
            "",  # empty string item (non-empty 'items' array, but an empty string element)
        ]},
    ]),
)

# -- relative image path resolved against the SPEC FILE's directory, not CWD
subdir = WORK / "subdir_for_cwd_test"
subdir.mkdir(exist_ok=True)
rel_img_spec_dir = WORK / "images_elsewhere"
rel_img_spec_dir.mkdir(exist_ok=True)
(rel_img_spec_dir / "pic.png").write_bytes(
    bytes.fromhex("89504e470d0a1a0a0000000d49484452000000010000000108020000009077"
                  "53de0000000c4944415408d763f8cfc0c0c0000003030100a8ee6cd80000000049454e44ae426082")
)
cwd_spec = {"meta": {}, "slides": [{"type": "image", "title": "T", "image": "pic.png"}]}
sp = write_spec("images_elsewhere/cwd_test.json", cwd_spec)
out_cwd = WORK / "cwd_test.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_cwd)], cwd=subdir)
check("relative image path resolves against spec dir, not CWD", p.returncode == 0 and out_cwd.is_file(),
      p.stderr.strip()[:300])

# -- --template pointing at a deck lacking the expected layouts -> clean error
bogus_pptx = WORK / "bogus_template.pptx"
blank_prs = Presentation()
blank_prs.save(str(bogus_pptx))
sp = write_spec("bad_template_spec.json", base_spec([{"type": "thanks"}]))
out_bt = WORK / "bad_template.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_bt),
         "--template", str(bogus_pptx)])
check("bad --template (missing expected layout) -> clean error, not AttributeError/traceback",
      p.returncode == 1 and "Traceback" not in p.stderr and "error:" in p.stderr,
      p.stderr.strip()[:300])

# -- --assets-dir override (an assets dir with no logo PNGs in it -- template
#    is passed separately since --assets-dir alone also controls the
#    *default* --template location, which is not what this case is testing).
empty_assets = WORK / "empty_assets"
empty_assets.mkdir(exist_ok=True)
sp = write_spec("assets_dir_spec.json", base_spec([{"type": "title"}, {"type": "back_cover"}]))
out_ad = WORK / "assets_dir_test.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_ad),
         "--assets-dir", str(empty_assets), "--template", str(TEMPLATE)])
check("--assets-dir override with no logos present does not crash", p.returncode == 0,
      p.stderr.strip()[:300])

# -- output dir creation (nested, non-existent)
nested_out = WORK / "a" / "b" / "c" / "deck.pptx"
sp = write_spec("nested_out_spec.json", base_spec([{"type": "thanks"}]))
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(nested_out)])
check("nested output directory auto-created", p.returncode == 0 and nested_out.is_file())

# -- overwrite behaviour: build twice to the same path
p1 = run([PY, str(BUILD), "--spec", str(sp), "--out", str(nested_out)])
check("re-building to the same --out overwrites cleanly", p1.returncode == 0 and nested_out.is_file())

# -- spec file not found / not valid JSON
p = run([PY, str(BUILD), "--spec", str(WORK / "nope.json"), "--out", str(WORK / "x.pptx")])
check("spec file not found -> clean exit 2", p.returncode == 2 and "Traceback" not in p.stderr)

bad_json = WORK / "bad.json"
bad_json.write_text("{not valid json", encoding="utf-8")
p = run([PY, str(BUILD), "--spec", str(bad_json), "--out", str(WORK / "x2.pptx")])
check("invalid JSON -> clean error, not traceback", p.returncode == 1 and "Traceback" not in p.stderr
      and "error:" in p.stderr)

# -- spec is a JSON array, not object
bad_top = WORK / "bad_top.json"
bad_top.write_text("[]", encoding="utf-8")
p = run([PY, str(BUILD), "--spec", str(bad_top), "--out", str(WORK / "x3.pptx")])
check("top-level JSON array -> clean error", p.returncode == 1 and "Traceback" not in p.stderr)


# -- overflow heuristic sanity check: a slide that clearly overflows its box
#    must WARN (stderr) and still exit 0; a short slide must NOT warn.
long_item = "Дуже довгий пункт списку, який явно переповнить блок, тому що складається з багатьох слів і речень. " * 4
overflow_spec = base_spec([{"type": "bullets", "title": "T", "items": [long_item] * 20}])
sp = write_spec("overflow_case.json", overflow_spec)
out_of = WORK / "overflow_case.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_of)])
check("overflow heuristic: 20 long items -> WARN + exit 0 (never hard-fails)",
      p.returncode == 0 and "WARN" in p.stderr and "do not fit" in p.stderr, p.stderr.strip()[:300])

short_spec = base_spec([{"type": "bullets", "title": "T", "items": ["Коротко.", "Ще коротше."]}])
sp = write_spec("no_overflow_case.json", short_spec)
out_no = WORK / "no_overflow_case.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_no)])
check("overflow heuristic: 2 short items -> no spurious WARN",
      p.returncode == 0 and "WARN" not in p.stderr, p.stderr.strip()[:300])


# =====================================================================
# 4. render_preview.py checks (needs soffice -- skipped cleanly if absent)
# =====================================================================
print("\n=== 4. render_preview.py ===")

if HAVE_SOFFICE:
    prev_out = WORK / "preview1"
    p = run([PY, str(PREVIEW), str(out1), "--out", str(prev_out), "--width", "800"])
    check("render_preview on example deck exits 0", p.returncode == 0, p.stderr[-500:])
    sheet = prev_out / "contact-sheet.png"
    check("contact sheet produced", sheet.is_file())

    import pymupdf as fitz  # noqa: E402

    png_count = len(list(prev_out.glob("slide-*.png")))
    check("PNG count == slide count", png_count == n_expected, f"{png_count} vs {n_expected}")

    # PDF page count == slide count (re-derive the PDF ourselves to check, since
    # render_preview.py deletes its tempdir -- re-run convert_to_pdf directly).
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        profile_dir = tmp_path / "lo_profile"
        profile_dir.mkdir()
        soffice = rp.find_soffice()
        pdf_path = rp.convert_to_pdf(out1, tmp_path, soffice, profile_dir)
        doc = fitz.open(str(pdf_path))
        check("PDF page count == slide count", doc.page_count == n_expected,
              f"{doc.page_count} vs {n_expected}")
        doc.close()
else:
    skip("render_preview on example deck exits 0", "soffice not on PATH")
    skip("contact sheet produced", "soffice not on PATH")
    skip("PNG count == slide count", "soffice not on PATH")
    skip("PDF page count == slide count", "soffice not on PATH")

# soffice missing -> clear message (monkeypatch find_soffice's lookups
# in-process). This does not need real soffice -- it simulates absence
# regardless of whether it's actually installed, so it always runs.
import shutil as _shutil
import pathlib as _pathlib
orig_is_file = _pathlib.Path.is_file
orig_which = _shutil.which
try:
    _pathlib.Path.is_file = lambda self: False
    _shutil.which = lambda cmd: None
    try:
        rp.find_soffice()
        check("find_soffice raises RuntimeError when soffice is unavailable", False, "no exception raised")
    except RuntimeError:
        check("find_soffice raises RuntimeError when soffice is unavailable", True)
finally:
    _pathlib.Path.is_file = orig_is_file
    _shutil.which = orig_which

# deck not found -> exit 2 (checked before find_soffice() is ever called, so
# this also always runs regardless of soffice availability).
p = run([PY, str(PREVIEW), str(WORK / "no-such-deck.pptx")])
check("render_preview: deck not found -> exit 2", p.returncode == 2 and "Traceback" not in p.stderr)

if HAVE_SOFFICE:
    # path with spaces + Cyrillic
    cyr_deck = WORK / "звіт 4 семестр.pptx"
    shutil.copy(out1, cyr_deck)
    cyr_out = WORK / "звіт-preview"
    p = run([PY, str(PREVIEW), str(cyr_deck), "--out", str(cyr_out)])
    check("render_preview handles Cyrillic + spaces in filename", p.returncode == 0 and
          (cyr_out / "contact-sheet.png").is_file(), p.stderr[-500:])

    # concurrent invocations (profile isolation) -- two decks rendered in parallel
    concurrent_specs = []
    for k in range(2):
        sp = write_spec(f"concurrent_{k}.json", base_spec([{"type": "thanks"}, {"type": "back_cover"}]))
        out_c = WORK / f"concurrent_{k}.pptx"
        rb = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_c)])
        concurrent_specs.append(out_c)
    procs = [
        subprocess.Popen([PY, str(PREVIEW), str(out_c), "--out", str(WORK / f"concurrent_prev_{k}")],
                          stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        for k, out_c in enumerate(concurrent_specs)
    ]
    rcs = [proc.wait(timeout=120) for proc in procs]
    errs = [proc.stderr.read() if proc.stderr else "" for proc in procs]
    check("two concurrent render_preview runs both succeed (profile isolation)",
          all(rc == 0 for rc in rcs), f"rcs={rcs} errs={errs}")
else:
    skip("render_preview handles Cyrillic + spaces in filename", "soffice not on PATH")
    skip("two concurrent render_preview runs both succeed (profile isolation)", "soffice not on PATH")


# =====================================================================
# 5. make_template.py checks
# =====================================================================
print("\n=== 5. make_template.py ===")

mt1 = WORK / "template_rebuilt.pptx"
p = run([PY, str(MAKE_TEMPLATE), "--source", str(TEMPLATE), "--out", str(mt1)])
check("make_template runs on the shipped template.pptx (idempotency source)", p.returncode == 0, p.stderr)

problems = check_ooxml_zip(mt1)
check("rebuilt template: zip-level OOXML sanity", not problems, "; ".join(problems[:5]))

# python-pptx can open it, and a deck can be built from it
prs_t = Presentation(str(mt1))
check("rebuilt template opens with python-pptx", True)
check("rebuilt template has zero slides", len(list(prs_t.slides)) == 0)
layout_names = {l.name for m in prs_t.slide_masters for l in m.slide_layouts}
check("rebuilt template still has the 2 layouts build_deck.py needs",
      bd.LAYOUT_CONTENT in layout_names and bd.LAYOUT_BACKCOVER in layout_names,
      f"{layout_names}")

sp = write_spec("via_rebuilt_template.json", base_spec([{"type": "title"}, {"type": "thanks"},
                                                         {"type": "back_cover"}]))
out_rt = WORK / "via_rebuilt_template.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_rt), "--template", str(mt1)])
check("build_deck.py works against the rebuilt template", p.returncode == 0, p.stderr)

# idempotency: re-run make_template on its OWN output (and again on the
# shipped assets/template.pptx, which already has zero slides)
mt2 = WORK / "template_rebuilt_twice.pptx"
p = run([PY, str(MAKE_TEMPLATE), "--source", str(mt1), "--out", str(mt2)])
check("make_template re-run on its own output succeeds (idempotent)", p.returncode == 0, p.stderr)
problems2 = check_ooxml_zip(mt2)
check("twice-rebuilt template: zip-level OOXML sanity", not problems2, "; ".join(problems2[:5]))

mt3 = WORK / "template_from_shipped.pptx"
p = run([PY, str(MAKE_TEMPLATE), "--source", str(TEMPLATE), "--out", str(mt3)])
check("make_template re-run directly on assets/template.pptx does not corrupt it",
      p.returncode == 0 and not check_ooxml_zip(mt3), p.stderr)
# and the ORIGINAL assets/template.pptx on disk must be untouched by any of this
check("original assets/template.pptx left untouched on disk", TEMPLATE.is_file() and not check_ooxml_zip(TEMPLATE))

# scrub check: no personal author name, no commentAuthors/comments/thumbnail/customXml parts
with zipfile.ZipFile(mt1) as z:
    names = z.namelist()
    core = z.read("docProps/core.xml").decode("utf-8") if "docProps/core.xml" in names else ""
    leaky = [n for n in names if ("comment" in n.lower() or "thumbnail" in n.lower()
                                   or n.lower().startswith("customxml/"))]
check("no comments/commentAuthors/thumbnail/customXml parts in rebuilt template", not leaky, f"{leaky}")
check("docProps/core.xml creator/lastModifiedBy scrubbed", "Vladyslav" not in core and "Holets" not in core)


# =====================================================================
# 6. meta.report_title (title-slide semester/report-title field)
# =====================================================================
print("\n=== 6. meta.report_title field ===")


def _title_slide_shape_tops(pptx_path: Path):
    """Returns a sorted list of (top_cm, left_cm, text[:60]) for every
    text-bearing shape on the deck's first (title) slide."""
    prs = Presentation(str(pptx_path))
    slide = prs.slides[0]
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append((round(sh.top / 360000, 3), round(sh.left / 360000, 3),
                        sh.text_frame.text.replace("\n", " | ")[:60]))
    return sorted(out)


rt_base_meta = {"dissertation_title": "T", "department": "D", "student_name": "S", "supervisor_name": "P"}

# -- absent: no report_title key at all -> geometry byte-for-byte matches
#    "no report_title support existed" (dissertation-topic block stays at
#    build_deck.TS_MAIN_T, no extra shape added, no empty reserved gap).
sp = write_spec("report_title_absent.json", {"meta": dict(rt_base_meta), "slides": [{"type": "title"}]})
out_absent = WORK / "report_title_absent.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_absent)])
check("report_title absent: builds clean (exit 0, no WARN)", p.returncode == 0 and "WARN" not in p.stderr, p.stderr)
tops_absent = _title_slide_shape_tops(out_absent)
main_tops_absent = [t for t, l, text in tops_absent if text == "T"]
check("report_title absent: dissertation-topic block at the unmodified TS_MAIN_T",
      main_tops_absent == [round(bd.TS_MAIN_T, 3)], f"{main_tops_absent} vs {bd.TS_MAIN_T}")
check("report_title absent: no extra report-title shape added",
      len(tops_absent) == 4, f"{len(tops_absent)} shapes: {tops_absent}")

# -- explicit None: must behave identically to absent (same positions, same
#    shape count), not render a blank/"None" line.
m_none = dict(rt_base_meta)
m_none["report_title"] = None
sp = write_spec("report_title_none.json", {"meta": m_none, "slides": [{"type": "title"}]})
out_none = WORK / "report_title_none.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_none)])
check("report_title None: builds clean (exit 0, no WARN)", p.returncode == 0 and "WARN" not in p.stderr, p.stderr)
check("report_title None: identical geometry to absent",
      _title_slide_shape_tops(out_none) == tops_absent, f"{_title_slide_shape_tops(out_none)} vs {tops_absent}")

# -- present (short, meta-level): renders a 5th shape at TS_MAIN_T, and
#    pushes the dissertation-topic block down by its height + the gap.
m_present = dict(rt_base_meta)
m_present["report_title"] = "Звіт за 4 семестр"
sp = write_spec("report_title_present.json", {"meta": m_present, "slides": [{"type": "title"}]})
out_present = WORK / "report_title_present.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_present)])
check("report_title present: builds clean (exit 0, no WARN)", p.returncode == 0 and "WARN" not in p.stderr, p.stderr)
tops_present = _title_slide_shape_tops(out_present)
check("report_title present: adds exactly one new shape", len(tops_present) == 5, f"{tops_present}")
rt_shapes = [(t, l, txt) for t, l, txt in tops_present if txt == "Звіт за 4 семестр"]
check("report_title present: rendered at the dissertation-topic block's old position",
      rt_shapes and rt_shapes[0][0] == round(bd.TS_MAIN_T, 3), f"{rt_shapes} vs TS_MAIN_T={bd.TS_MAIN_T}")
main_tops_present = [t for t, l, text in tops_present if text == "T"]
expected_main_top = round(bd.TS_MAIN_T + bd.TS_REPORT_TITLE_H + bd.TS_REPORT_TITLE_GAP_CM, 3)
check("report_title present: dissertation-topic block pushed down by report-title height + gap",
      main_tops_present == [expected_main_top], f"{main_tops_present} vs {expected_main_top}")

# -- non-string (a bare number): must coerce to its string form, not crash.
m_num = dict(rt_base_meta)
m_num["report_title"] = 4
sp = write_spec("report_title_nonstring.json", {"meta": m_num, "slides": [{"type": "title"}]})
out_num = WORK / "report_title_nonstring.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_num)])
check("report_title non-string (int): builds clean, does not crash", p.returncode == 0, p.stderr)
tops_num = _title_slide_shape_tops(out_num)
check("report_title non-string (int): coerced to its str() form ('4')",
      any(text == "4" for _, _, text in tops_num), f"{tops_num}")

# -- per-slide override: slide's own `report_title` wins over `meta`'s
#    (same precedent as every other title field), including overriding a
#    meta-level value to a different per-slide one.
sp = write_spec("report_title_override.json", {
    "meta": m_present,
    "slides": [{"type": "title", "report_title": "Override за 5 семестр"}],
})
out_override = WORK / "report_title_override.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_override)])
check("report_title per-slide override: builds clean", p.returncode == 0, p.stderr)
tops_override = _title_slide_shape_tops(out_override)
check("report_title per-slide override: slide's own value wins over meta's",
      any(text == "Override за 5 семестр" for _, _, text in tops_override)
      and not any(text == "Звіт за 4 семестр" for _, _, text in tops_override),
      f"{tops_override}")

# -- unknown-meta-key WARN must NOT fire for report_title itself (it's a
#    known key now, not a typo).
sp = write_spec("report_title_known_key.json", {"meta": m_present, "slides": [{"type": "title"}]})
out_kk = WORK / "report_title_known_key.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_kk)])
check("'report_title' is a known meta/title key (no 'unrecognised field' WARN)",
      p.returncode == 0 and "unrecognised field" not in p.stderr.lower(), p.stderr)


# =====================================================================
# 7. definition-slide overflow WARN names the specific block
# =====================================================================
print("\n=== 7. definition overflow WARN block naming ===")

long_goal = ("Дуже довгий текст мети дослідження, який явно переповнить блок, тому що "
             "містить багато слів і речень про щось дуже важливе. " * 3)
def_spec = {
    "meta": dict(rt_base_meta),
    "slides": [{
        "type": "definition",
        "object": {"text": "Короткий об'єкт."},
        "subject": {"text": "Короткий предмет."},
        "goal": {"text": long_goal},
    }],
}
sp = write_spec("definition_overflow_goal.json", def_spec)
out_def = WORK / "definition_overflow_goal.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_def)])
check("definition overflow (goal block): builds anyway (exit 0)", p.returncode == 0, p.stderr)
check("definition overflow WARN names the specific overflowing block ('Мета дослідження')",
      "WARN" in p.stderr and "Мета дослідження" in p.stderr, p.stderr.strip()[:300])
check("definition overflow WARN does NOT also blame the short, fitting blocks",
      "Об'єкт дослідження" not in p.stderr and "Предмет дослідження" not in p.stderr, p.stderr.strip()[:300])
check("definition overflow WARN states estimated vs. available line counts",
      "estimated ~" in p.stderr and "lines fit" in p.stderr, p.stderr.strip()[:300])


# =====================================================================
# 8. Fidelity fixes: bold body/bullets/conclusions/references, sizes,
#    and the recalibrated bold-aware overflow heuristic
# =====================================================================
print("\n=== 8. Bold-text fidelity fixes + overflow recalibration ===")

check("CONCLUSIONS_SIZE_PT matches the measured original (20pt, was 18pt)",
      bd.CONCLUSIONS_SIZE_PT == 20, f"{bd.CONCLUSIONS_SIZE_PT}")
check("TS_WORDMARK_PT matches the measured original (19pt, was 16pt)",
      bd.TS_WORDMARK_PT == 19, f"{bd.TS_WORDMARK_PT}")

fidelity_spec = {
    "meta": dict(rt_base_meta),
    "slides": [
        {"type": "bullets", "title": "T", "items": ["Пункт один.", "Пункт два."]},
        {"type": "conclusions", "title": "T", "items": ["Висновок один."]},
        {"type": "references", "title": "T", "items": ["Автор, І. (2024). Назва. Журнал, 1, 1-2."]},
    ],
}
sp = write_spec("fidelity_bold_check.json", fidelity_spec)
out_fb = WORK / "fidelity_bold_check.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_fb)])
check("bold-fidelity spec builds clean", p.returncode == 0 and "WARN" not in p.stderr, p.stderr)
if out_fb.is_file():
    prs_fb = Presentation(str(out_fb))
    bold_flags = []
    for slide in prs_fb.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run_ in para.runs:
                    if run_.text.strip() in ("Пункт один.", "Пункт два.", "Висновок один.") or \
                       run_.text.startswith("Автор, І."):
                        bold_flags.append((run_.text, run_.font.bold))
    check("bullets/conclusions/references items are rendered bold (b=1), matching the original",
          bold_flags and all(b is True for _, b in bold_flags), f"{bold_flags}")

# -- regression test for the exact bug this recalibration fixes: a 5-item
#    numbered list whose items are each ~50-57 chars (i.e. squarely in the
#    "looked fine at the old 0.52/0.56 factors, but the real bold-rendered
#    text actually wraps to 2 lines and creeps into the footer" range found
#    while rebuilding assets/example-spec.json) must now WARN.
borderline_items = [
    "Дослідити існуючі методи та моделі за темою дослідження;",  # 56 chars
    "Розробити новий показник/критерій ефективності;",            # 47 chars
    "Дослідити методи збору та обробки даних експерименту;",      # 53 chars
    "Розробити власний метод розв'язання поставленої задачі;",    # 55 chars
    "Розробити програмне забезпечення для розробленого методу.",  # 57 chars
]
borderline_spec = {"meta": dict(rt_base_meta),
                    "slides": [{"type": "bullets", "title": "T", "items": borderline_items}]}
sp = write_spec("borderline_bold_overflow.json", borderline_spec)
out_bl = WORK / "borderline_bold_overflow.pptx"
p = run([PY, str(BUILD), "--spec", str(sp), "--out", str(out_bl)])
check("bold overflow calibration: 50-57 char items (fit under the old regular-text factor, "
      "genuinely wrap to 2 lines once bold) now correctly WARN",
      p.returncode == 0 and "WARN" in p.stderr, p.stderr.strip()[:300])


# =====================================================================
# Summary
# =====================================================================
print("\n" + "=" * 70)
n_pass = sum(1 for _, ok, _ in RESULTS if ok)
n_fail = len(RESULTS) - n_pass
n_skip = len(SKIPPED)
print(f"SUMMARY: {n_pass}/{len(RESULTS)} passed, {n_fail} failed, {n_skip} skipped")
if n_fail:
    print("\nFAILED:")
    for name, ok, detail in RESULTS:
        if not ok:
            print(f"  - {name}: {detail}")
    sys.exit(1)
sys.exit(0)
