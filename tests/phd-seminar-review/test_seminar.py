# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "python-pptx>=1.0",
#     "python-docx>=1.1",
#     "pymupdf>=1.24",
#     "pillow",
# ]
# ///
"""Plain-assert (no pytest) regression harness for the phd-seminar-review skill
scripts: check_deck.py, extract_deck.py, render_preview.py.

    uv run tests/phd-seminar-review/test_seminar.py

Works from any CWD; every deck is generated with python-pptx into a temporary
directory, so no real (personal) presentation is needed. Covers: a complete
121 seminar deck that must pass without errors, every major rule on
deliberately broken variants (missing architecture result, other specialty,
object = subject, placeholders, calques, foreign publications, missing
sections), cross-checking against a dissertation .docx, extraction to
deck.md/deck.json, error handling, and (with LibreOffice installed) the PDF
path and the preview renderer. Prints a SUMMARY line and exits non-zero on
failure.
"""
from __future__ import annotations

import atexit
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = REPO_ROOT / "skills" / "phd-seminar-review" / "scripts"
CHECK, EXTRACT, RENDER = SCRIPTS / "check_deck.py", SCRIPTS / "extract_deck.py", SCRIPTS / "render_preview.py"
TMP = Path(tempfile.mkdtemp(prefix="seminar-test-"))
atexit.register(shutil.rmtree, TMP, ignore_errors=True)
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice") or (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists() else None)

RESULTS = {"pass": 0, "fail": 0, "skip": 0}

TOPIC = "Методи та засоби оркестрації мікросервісів у хмарних енергетичних інформаційних системах"
OBJECT = "Процес розгортання та функціонування мікросервісних програмних систем у хмарних середовищах."
SUBJECT = "Моделі, методи та архітектурні засоби оркестрації мікросервісів у хмарних програмних системах."
AIM = ("Мета дослідження: зменшення часу відновлення мікросервісних програмних систем шляхом розроблення "
       "методів та архітектури оркестрації.")
TASKS = [
    "1. Провести аналіз методів оркестрації мікросервісів у хмарних середовищах.",
    "2. Розробити модель залежностей мікросервісів програмної системи.",
    "3. Удосконалити метод планування розгортання мікросервісів.",
    "4. Розробити архітектуру програмної системи оркестрації та перевірити її експериментально.",
]
NOVELTY = [
    "1. Вперше запропоновано модель залежностей мікросервісів програмної системи, яка відрізняється урахуванням "
    "часових вікон відмов, що дозволило прогнозувати каскадні збої.",
    "2. Удосконалено метод планування розгортання мікросервісів, який відрізняється використанням графа "
    "залежностей, що дозволило зменшити час відновлення на 27 %.",
    "3. Удосконалено архітектуру програмної системи оркестрації, яка відрізняється виділенням компонента "
    "прогнозування відмов, що дозволило скоротити простої в 1,8 раза.",
]
CONCLUSIONS = [
    "1. Аналіз показав, що відомі методи оркестрації не враховують часові вікна відмов.",
    "2. Розроблена модель залежностей дозволяє прогнозувати каскадні збої з точністю 91 %.",
    "3. Удосконалений метод планування зменшив час відновлення на 27 %.",
    "4. Архітектура системи оркестрації скоротила простої в 1,8 раза; її впроваджено в ТОВ «Приклад».",
]
FUTURE = "Подальші дослідження: адаптивне масштабування оркестратора."
PUBS = [
    "За темою дисертації опубліковано 4 статті (з них 1 у виданні Scopus Q3) та 3 тези конференцій.",
    "1. Петренко О.І., Іваненко П.П. Модель залежностей мікросервісів. Адаптивні системи автоматичного управління. "
    "2025. № 1 (46). С. 10–18.",
    "2. Петренко О.І. Метод планування розгортання. Технічна інженерія. 2025. № 2. С. 30–37.",
    "3. Petrenko O. Orchestration architecture. Energy Informatics (Scopus Q3). 2026. Vol. 9. P. 1–12.",
    "4. Петренко О.І., Іваненко П.П. Прогнозування відмов. Вісник КПІ. 2026. № 1. С. 5–11.",
]


def make_deck(path: Path, *, novelty=None, specialty="121 «Інженерія програмного забезпечення»", subject=SUBJECT,
              pubs=None, extra_text=None, drop=(), arch_slide=True, notes=True, no_future=False) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide(title, paras, size=20, picture=False):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.9))
        tb.text_frame.text = title.upper()
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        body = s.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(12), Inches(5.5))
        tf = body.text_frame
        for k, p in enumerate(paras):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            para.text = p
            for r in para.runs:
                r.font.size = Pt(size)
        num = s.shapes.add_textbox(Inches(12.6), Inches(7.0), Inches(0.6), Inches(0.4))
        num.text_frame.text = str(len(prs.slides))
        if picture:
            for k in range(7):
                s.shapes.add_shape(1, Inches(1 + k), Inches(5), Inches(0.8), Inches(0.5))
        if notes:
            s.notes_slide.notes_text_frame.text = "Текст доповіді."
        return s

    title = prs.slides.add_slide(blank)
    tb = title.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(6.5))
    tf = tb.text_frame
    for k, p in enumerate([
        "Національний технічний університет України «Київський політехнічний інститут імені Ігоря Сікорського»",
        "ПЕТРЕНКО Олена Іванівна", TOPIC.upper(), f"Спеціальність {specialty}",
        "Дисертація на здобуття ступеня доктора філософії", "Науковий керівник",
        "доктор технічних наук, професор ІВАНЕНКО Петро Петрович", "Київ – 2026"]):
        (tf.paragraphs[0] if k == 0 else tf.add_paragraph()).text = p
    sections = [
        ("relevance", "Актуальність теми дослідження",
         ["Простої хмарних енергетичних систем коштують операторам до 5 % річного доходу.",
          "Відомі оркестратори не враховують каскадні відмови мікросервісів."]),
        ("osa", "Об'єкт, предмет та мета дослідження",
         [f"Об'єкт дослідження: {OBJECT}", f"Предмет дослідження: {subject}", AIM,
          "Методи дослідження: теорія графів, імітаційне моделювання, методи програмної інженерії."]),
        ("tasks", "Завдання дослідження", TASKS),
        ("novelty", "Наукова новизна отриманих результатів", novelty if novelty is not None else NOVELTY),
        ("r1", "Модель залежностей мікросервісів", ["Граф G = (V, E), де V — мікросервіси програмної системи."]),
        ("arch", "Архітектура програмної системи оркестрації" if arch_slide else "Програмна реалізація",
         ["Компоненти: планувальник, прогнозувач відмов, агент розгортання."]),
        ("exp", "Результати експериментів", ["Час відновлення зменшено на 27 % порівняно з Kubernetes."]),
        ("practical", "Практичне значення", ["Акт впровадження ТОВ «Приклад» від 12.03.2026."]),
        ("conclusions", "Висновки", CONCLUSIONS + ([] if no_future else [FUTURE])),
        ("publications", "Публікації та апробація результатів",
         (pubs if pubs is not None else PUBS) + ["Апробація: 3 міжнародні конференції."]),
        ("thanks", "Дякую за увагу!", []),
    ]
    for key, t, paras in sections:
        if key in drop:
            continue
        slide(t, paras, picture=(key == "arch"))
    if extra_text:
        for page in (extra_text if isinstance(extra_text[0], list) else [extra_text]):
            slide("Додатково", page)
    prs.save(str(path))
    return path


def run(script: Path, *args) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(script), *map(str, args)], capture_output=True, text=True)


def check(name: str, cond: bool, detail: str = ""):
    if cond:
        RESULTS["pass"] += 1
        print(f"[PASS] {name}")
    else:
        RESULTS["fail"] += 1
        print(f"[FAIL] {name}\n{detail[-3000:]}")


def skip(name: str, why: str):
    RESULTS["skip"] += 1
    print(f"[SKIP] {name}: {why}")


def codes(out: str, level: str | None = None) -> set[str]:
    found = set()
    for line in out.splitlines():
        parts = line.split()
        if len(parts) >= 2 and parts[0] in ("ERROR", "WARN", "INFO") and (level is None or parts[0] == level):
            found.add(parts[1].rstrip(":"))
    return found


def main() -> int:
    good = make_deck(TMP / "good.pptx")
    r = run(CHECK, good, "--json", TMP / "good.json")
    check("good deck: exit 0, no errors", r.returncode == 0 and not codes(r.stdout, "ERROR"), r.stdout + r.stderr)
    data = json.loads((TMP / "good.json").read_text(encoding="utf-8"))
    check("good deck: candidate surname from the title slide", data["title"]["candidate"] == "Петренко",
          str(data["title"]))
    check("good deck: topic found", data["title"]["topic"].lower() == TOPIC.lower(), data["title"]["topic"])
    check("good deck: every required section found",
          all(data["sections"][k] for k in ("relevance", "object", "subject", "aim", "tasks", "novelty",
                                            "conclusions", "publications", "thanks")), str(data["sections"]))
    check("good deck: 4 tasks / 3 novelty / 4 conclusions",
          (len(data["tasks"]), len(data["novelty"]), len(data["conclusions"])) == (4, 3, 4),
          json.dumps({k: data[k] for k in ("tasks", "novelty", "conclusions")}, ensure_ascii=False))
    check("good deck: 4 publication entries, no authorship warnings",
          len(data["publications"]["entries"]) == 4 and "P03" not in codes(r.stdout), r.stdout)
    check("good deck: no architecture warnings", not codes(r.stdout) & {"A01", "A02", "A03", "A04"}, r.stdout)

    no_arch = make_deck(TMP / "no_arch.pptx", novelty=NOVELTY[:2] + [
        "3. Вперше розроблено програмне забезпечення системи оркестрації, яке відрізняється модульністю, що "
        "дозволило скоротити простої."], arch_slide=False)
    r = run(CHECK, no_arch)
    check("no architecture result: A01 ERROR, exit 1", "A01" in codes(r.stdout, "ERROR") and r.returncode == 1,
          r.stdout)
    check("no architecture slide: A03 WARN", "A03" in codes(r.stdout, "WARN"), r.stdout)

    r = run(CHECK, make_deck(TMP / "other.pptx", specialty="122 «Комп'ютерні науки»"))
    check("other specialty: T02 ERROR", "T02" in codes(r.stdout, "ERROR"), r.stdout)

    r = run(CHECK, make_deck(TMP / "same.pptx", subject=OBJECT))
    check("subject equals object: O01 ERROR", "O01" in codes(r.stdout, "ERROR"), r.stdout)

    r = run(CHECK, make_deck(TMP / "ph.pptx", extra_text=["TODO: вставити графік", "Lorem ipsum dolor sit amet"]))
    check("placeholder text: F05 ERROR", "F05" in codes(r.stdout, "ERROR"), r.stdout)

    r = run(CHECK, make_deck(TMP / "lang.pptx", extra_text=[
        "Даний метод являється ефективним, так як об'єм даних зростає на протязі року на 10 процентів."]))
    check("calques: L01 for являється/об'єм/на протязі", "L01" in codes(r.stdout, "WARN")
          and all(w in r.stdout for w in ("являється", "об'єм", "на протязі")), r.stdout)

    r = run(CHECK, make_deck(TMP / "lang2.pptx", extra_text=[
        ["Даний метод перевірено на кластері з п'яти вузлів у різних конфігураціях навантаження.",
         "Рис. 2.1. Схема експерименту з кластером і точність 0.95 у першій серії вимірювань."],
        ["Даний метод перевірено повторно на тому ж кластері з іншими конфігураціями вузлів.",
         "Рис. 2.1. Повторна схема; DIAG=1 означає відмову вузла під час тестування кластера."]]))
    l01 = [l for l in r.stdout.splitlines() if " L01 " in l and "даний" in l]
    check("L01 reported once per calque with the other slides listed",
          len(l01) == 1 and "також слайди 14" in l01[0], r.stdout)
    check("F08 duplicated figure number; L03 mixed decimals; no L02 for DIAG=",
          "F08" in codes(r.stdout, "WARN") and "L03" in codes(r.stdout, "INFO") and "DIAG" not in r.stdout, r.stdout)
    r = run(CHECK, make_deck(TMP / "nofuture.pptx", no_future=True))
    check("no future-research line in conclusions: C01", "C01" in codes(r.stdout, "WARN"), r.stdout)

    r = run(CHECK, make_deck(TMP / "pubs.pptx", pubs=PUBS[:2] + [
        "2. Сидоренко В.В. Чужа стаття про бази даних. Зв'язок. 2022. № 3. С. 31–34."]))
    check("publication without the candidate: P03, too few: P02",
          {"P03", "P02"} <= codes(r.stdout, "WARN"), r.stdout)

    data_pub = json.loads((TMP / "good.json").read_text(encoding="utf-8"))["publications"]
    check("publication credits: Q3 = 2, 2-author фахова = 1, total 5",
          [c["credit"] for c in data_pub["credits"]] == [1.0, 1.0, 2.0, 1.0] and data_pub["credit_total"] == 5,
          json.dumps(data_pub["credits"], ensure_ascii=False))
    r = run(CHECK, make_deck(TMP / "credits.pptx", pubs=[
        "1. Петренко О.І., Іваненко П.П., Сидоренко В.В. Перша стаття. Вісник КПІ. 2025. № 1. С. 5–11.",
        "2. Петренко О.І., Іваненко П.П., Коваль О.В. Друга стаття. Вісник КПІ. 2025. № 1. С. 12–19.",
        "3. Петренко О.І. Тези. Матеріали міжнародної науково-практичної конференції. Київ, 2025. С. 3–4.",
        "4. Петренко О.І. Стаття. Технічна інженерія. 2026. № 2. С. 30–37."]))
    check("> 2 authors = 0.5, theses = 0 -> P02; same issue twice -> P08",
          {"P02", "P08"} <= codes(r.stdout, "WARN") and "разом 2" in r.stdout, r.stdout)
    r = run(CHECK, no_arch)
    check("software as novelty without architecture: N06", "N06" in codes(r.stdout, "WARN"), r.stdout)

    r = run(CHECK, make_deck(TMP / "nogr.pptx", novelty=[
        "1. Розроблено модель залежностей мікросервісів програмної системи для прогнозування збоїв у хмарі.",
        "2. Архітектура системи оркестрації мікросервісів програмної системи з компонентом прогнозування."]))
    check("novelty without gradation/distinction/effect: N01, N02, N03",
          {"N01", "N02", "N03"} <= codes(r.stdout, "WARN"), r.stdout)

    r = run(CHECK, make_deck(TMP / "missing.pptx", drop=("relevance", "conclusions")))
    check("missing relevance and conclusions: S01 ERROR", "S01" in codes(r.stdout, "ERROR")
          and "Актуальність" in r.stdout and "Висновки" in r.stdout, r.stdout)

    r = run(CHECK, good, "--minutes", "5")
    check("too many slides for the time: F01", "F01" in codes(r.stdout, "WARN"), r.stdout)

    # Cross-check with the dissertation.
    diss = Document()
    for p in [f"Тема: {TOPIC}", "Об'єкт дослідження – " + OBJECT, "Предмет дослідження – " + SUBJECT,
              "Наукова новизна одержаних результатів:", *[n[3:] for n in NOVELTY], "Практичне значення …"]:
        diss.add_paragraph(p)
    diss.save(str(TMP / "diss.docx"))
    r = run(CHECK, good, "--dissertation", TMP / "diss.docx")
    check("matching dissertation: no X01/X02/X03", not codes(r.stdout) & {"X01", "X02", "X03"}, r.stdout)
    other = Document()
    for p in ["Тема: Зовсім інша тема про нейронні мережі для прогнозування погоди",
              "Об'єкт дослідження – процес навчання нейронних мереж.", "Наукова новизна:", "1. Інше.",
              "Практичне значення"]:
        other.add_paragraph(p)
    other.save(str(TMP / "other.docx"))
    r = run(CHECK, good, "--dissertation", TMP / "other.docx")
    check("different dissertation: X01 ERROR, X02, X03", "X01" in codes(r.stdout, "ERROR")
          and {"X02", "X03"} <= codes(r.stdout, "WARN"), r.stdout)

    # Extraction and error handling.
    r = run(EXTRACT, good, "--out", TMP / "rev")
    md = (TMP / "rev" / "deck.md")
    check("extract_deck writes deck.md/deck.json", r.returncode == 0 and md.exists()
          and "## Слайд 5: НАУКОВА НОВИЗНА" in md.read_text(encoding="utf-8")
          and (TMP / "rev" / "deck.json").exists(), r.stdout + r.stderr)
    (TMP / "x.txt").write_text("hi", encoding="utf-8")
    r = run(CHECK, TMP / "x.txt")
    check("unsupported format: exit 2 with a clean error", r.returncode == 2 and r.stderr.startswith("error:"),
          r.stdout + r.stderr)
    r = run(CHECK, TMP / "missing.pdf")
    check("missing file: exit 2", r.returncode == 2, r.stderr)

    if not SOFFICE:
        skip("PDF input and preview", "LibreOffice not installed")
    else:
        r = run(RENDER, good, "--out", TMP / "prev")
        pdf = TMP / "prev" / "good.pdf"
        check("render_preview: PNG per slide + contact sheet",
              r.returncode == 0 and (TMP / "prev" / "contact-sheet.png").exists()
              and len(list((TMP / "prev").glob("slide-*.png"))) == 12, r.stdout + r.stderr)
        r = run(CHECK, pdf, "--json", TMP / "pdf.json")
        d = json.loads((TMP / "pdf.json").read_text(encoding="utf-8")) if (TMP / "pdf.json").exists() else {}
        check("PDF input: no errors and the same structure", r.returncode == 0 and d
              and all(d["sections"][k] for k in ("relevance", "object", "tasks", "novelty", "conclusions"))
              and len(d["novelty"]) == 3, r.stdout + r.stderr)

    print(f"\nSUMMARY: {RESULTS['pass']} passed, {RESULTS['fail']} failed, {RESULTS['skip']} skipped")
    return 1 if RESULTS["fail"] else 0


if __name__ == "__main__":
    sys.exit(main())
