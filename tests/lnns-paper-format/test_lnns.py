# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "python-docx>=1.1",
#     "pyyaml>=6",
#     "pypandoc_binary>=1.13",
#     "pymupdf>=1.24",
#     "pillow",
# ]
# ///
"""Plain-assert (no pytest) regression harness for the lnns-paper-format skill
scripts: build_paper.py, check_paper.py, renumber_refs.py, import_source.py,
render_preview.py.

    uv run tests/lnns-paper-format/test_lnns.py

Works from any CWD; every path is derived from this file's location or a fresh
temporary directory. The interpreter running this file has every dependency
the skill scripts declare, so it runs them directly (sys.executable).

Covers: the shipped example end-to-end (Word layout against the Springer
splnproc template, the llncs LaTeX project), all-or-nothing build errors,
every major check_paper.py rule on deliberately broken sources, reference
renumbering, the importer on the built .docx and .tex, zip-level OOXML
well-formedness, and (when LibreOffice is installed) rendering plus the PDF
checks. A TeX engine, if installed, also compiles the LaTeX project. Prints a
final SUMMARY line and exits non-zero if anything failed.
"""
from __future__ import annotations

import atexit
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import yaml
from docx import Document
from docx.oxml.ns import qn

REPO_ROOT = Path(__file__).resolve().parents[2]
SKILL_DIR = REPO_ROOT / "skills" / "lnns-paper-format"
SCRIPTS = SKILL_DIR / "scripts"
BUILD = SCRIPTS / "build_paper.py"
CHECK = SCRIPTS / "check_paper.py"
RENUM = SCRIPTS / "renumber_refs.py"
IMPORT = SCRIPTS / "import_source.py"
PREVIEW = SCRIPTS / "render_preview.py"
EXAMPLE = SKILL_DIR / "assets" / "example" / "paper.md"
PY = sys.executable

_tmp = tempfile.TemporaryDirectory(prefix="lnns-paper-format-test-")
atexit.register(_tmp.cleanup)
WORK = Path(_tmp.name)

sys.path.insert(0, str(SCRIPTS))
import lnns_common as lc  # noqa: E402

RESULTS: list[tuple[str, bool, str]] = []
SKIPPED: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    RESULTS.append((name, bool(cond), detail))
    print(f"[{'PASS' if cond else 'FAIL'}] {name}" + (f" -- {detail}" if detail and not cond else ""))


def skip(name: str, reason: str) -> None:
    SKIPPED.append(name)
    print(f"[SKIP] {name} -- {reason}")


def run(script: Path, *args, cwd: Path | None = None) -> subprocess.CompletedProcess:
    return subprocess.run([PY, str(script), *map(str, args)], capture_output=True, text=True, cwd=cwd)


def workdir(name: str) -> Path:
    d = WORK / name
    d.mkdir(parents=True, exist_ok=True)
    for f in (SKILL_DIR / "assets" / "example").iterdir():
        shutil.copy(f, d / f.name)
    return d


def split_source(path: Path) -> tuple[dict, str]:
    src = lc.load_source(path)
    return src.meta, src.body


def write_source(path: Path, meta: dict, body: str) -> None:
    path.write_text("---\n" + yaml.safe_dump(meta, allow_unicode=True, sort_keys=False) + "---\n" + body,
                    encoding="utf-8")


def codes(output: str) -> set[str]:
    return {m.group(2) for m in re.finditer(r"(?m)^(ERROR|WARN|INFO)\s+(\S+):", output)}


def level_codes(output: str, level: str) -> set[str]:
    return {m.group(1) for m in re.finditer(rf"(?m)^{level}\s+(\S+):", output)}


# ---------------------------------------------------------------- build: example

d = workdir("example")
p = run(BUILD, d / "paper.md", "--docx", d / "paper.docx", "--latex", d / "latex", "--verbose")
check("example builds (docx + latex)", p.returncode == 0, p.stderr[-800:])
check("verbose log lists sections/equations", "section 1 " in p.stderr and "equation (2)" in p.stderr, p.stderr[-400:])

doc = Document(str(d / "paper.docx"))
s = doc.sections[0]
check("A4 page", abs(s.page_width.mm - 210) < 0.5 and abs(s.page_height.mm - 297) < 0.5)
tw = s.page_width.mm - s.left_margin.mm - s.right_margin.mm
th = s.page_height.mm - s.top_margin.mm - s.bottom_margin.mm
check("text area 122 x 193 mm", abs(tw - 122) < 0.5 and abs(th - 193) < 0.5, f"{tw:.1f} x {th:.1f}")
check("no header/footer text", not any(p.text.strip() for p in s.header.paragraphs + s.footer.paragraphs))

paras = doc.paragraphs
texts = [p.text for p in paras]
title = paras[0]
check("title first, 14 pt bold centred",
      title.text.startswith("Drift-Aware") and all(r.bold for r in title.runs if r.text.strip())
      and all(r.font.size.pt == 14 for r in title.runs if r.text.strip())
      and title.alignment == 1, f"{title.text!r}")
authors = paras[1]
sup = [r.text for r in authors.runs if r.font.superscript]
check("author line has affiliation superscripts and ORCIDs",
      "1" in sup and "1,2" in sup and any(t.startswith("[0000-") for t in sup), str(sup))
check("corresponding author marked with envelope", "✉" in authors.text)
check("affiliation 9 pt with e-mails",
      any("Kyiv 03056" in t and "i.petrenko@example.com" in t for t in texts))
abs_p = next(p for p in paras if p.text.startswith("Abstract."))
check("abstract 9 pt, indented 1 cm both sides",
      all(r.font.size.pt == 9 for r in abs_p.runs if r.text.strip())
      and abs(abs_p.paragraph_format.left_indent.mm - 10) < 0.2
      and abs(abs_p.paragraph_format.right_indent.mm - 10) < 0.2)
check("abstract label bold run-in", abs_p.runs[0].bold and abs_p.runs[0].text.strip() == "Abstract.")
kw = next(p for p in paras if p.text.startswith("Keywords:"))
check("keywords comma-separated, final period", kw.text.endswith("Adaptive Control.") and ", " in kw.text)

heads = [t for t in texts if re.match(r"^\d+(\.\d+)?\t", t)]
check("level 1/2 headings numbered 1..5 and 3.1/3.2",
      [h.split("\t")[0] for h in heads] == ["1", "2", "3", "3.1", "3.2", "4", "5"], str(heads))
h1 = next(p for p in paras if p.text.startswith("1\tIntroduction"))
check("level 1 heading 12 pt bold", all(r.bold and r.font.size.pt == 12 for r in h1.runs if r.text.strip()))
first_after = paras[paras.index(h1) + 1]
check("first paragraph after heading not indented", (first_after.paragraph_format.first_line_indent or 0) == 0)
second = paras[paras.index(h1) + 2]
check("following paragraph indented 0.4 cm",
      second.paragraph_format.first_line_indent is not None
      and abs(second.paragraph_format.first_line_indent.mm - 4.0) < 0.1)
runin = next(p for p in paras if p.text.startswith("Acceptance Test."))
check("level 3 heading is a bold run-in", runin.runs[0].bold and runin.runs[0].text == "Acceptance Test."
      and "guarantees" in runin.text)

body_xml = doc.element.body
m_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
eq_paras = [p for p in paras if p._p.find(f"{m_ns}oMath") is not None and re.search(r"\(\d+\)$", p.text)]
check("two numbered display equations", [re.search(r"\((\d+)\)$", p.text).group(1) for p in eq_paras] == ["1", "2"])

figcap = next(p for p in paras if p.text.startswith("Fig. 1."))
check("figure caption: bold upright label, 9 pt, below figure",
      figcap.runs[0].bold and figcap.runs[0].italic is False and figcap.runs[0].font.size.pt == 9
      and paras[paras.index(figcap) - 1]._p.find(f".//{qn('w:drawing')}") is not None)
tabcap = next(p for p in paras if p.text.startswith("Table 1."))
nxt = tabcap._p.getnext()
check("table caption above the table", nxt is not None and nxt.tag == qn("w:tbl"))
tbl = doc.tables[0]
grid = sum(int(c.get(qn("w:w"))) for c in tbl._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")))
check("table fits the 122 mm text width", grid <= round(122 * 1440 / 25.4) + 2, f"{grid} twips")
borders = tbl._tbl.tblPr.find(qn("w:tblBorders"))
check("table has horizontal rules only",
      borders.find(qn("w:top")).get(qn("w:val")) == "single"
      and borders.find(qn("w:insideV")).get(qn("w:val")) == "nil")

check("credits before references",
      texts.index(next(t for t in texts if t.startswith("Disclosure of Interests.")))
      < texts.index("References"))
refs = [t for t in texts[texts.index("References") + 1:] if re.match(r"^\d+\.\t", t)]
check("13 numbered references", len(refs) == 13, str(len(refs)))
ref_p = paras[texts.index("References") + 1]
check("references 9 pt", all(r.font.size.pt == 9 for r in ref_p.runs if r.text.strip()))
check("citation range normalised", any("[5–7]" in t for t in texts))

with zipfile.ZipFile(d / "paper.docx") as z:
    bad = []
    for n in z.namelist():
        if n.endswith(".xml") or n.endswith(".rels"):
            try:
                ET.fromstring(z.read(n))
            except ET.ParseError as exc:
                bad.append(f"{n}: {exc}")
    check("docx XML well-formed", not bad, "; ".join(bad))
    check("no macros in docx", not any("vbaProject" in n for n in z.namelist()))
    docpr = re.findall(r'<wp:docPr [^>]*descr="([^"]*)"', z.read("word/document.xml").decode())
    check("figure alt text in Word image description", any(d_.startswith("Line chart") for d_ in docpr), str(docpr))

# ---------------------------------------------------------------- build: LaTeX

tex = (d / "latex" / "paper.tex").read_text(encoding="utf-8")
check("latex files present", all((d / "latex" / f).is_file() for f in ("paper.tex", "llncs.cls", "splncs04.bst", "fig1.png")))
check("llncs runningheads class", r"\documentclass[runningheads]{llncs}" in tex)
check("authors with \\inst, \\orcidID and \\and",
      r"Ivan Petrenko\inst{1}" in tex and r"\orcidID{0000-0000-0000-0001}" in tex and r" \and" in tex)
check("\\titlerunning and \\authorrunning", r"\titlerunning{" in tex and r"\authorrunning{I. Petrenko and O. Kovalenko}" in tex)
check("\\institute with \\email", r"\institute{" in tex and r"\email{i.petrenko@example.com, o.kovalenko@example.com}" in tex)
check("keywords with \\and", r"\keywords{Concept Drift \and Gradient Boosting" in tex)
check("numbered equation environments", tex.count(r"\begin{equation}") == 2)
check("\\cite keys", r"\cite{ref5,ref6,ref7}" in tex and r"\cite{ref13}" in tex)
check("figure float with Description and caption",
      r"\Description{Line chart" in tex and r"\caption{Validation error" in tex and r"\label{fig1}" in tex)
check("table float caption before tabular", tex.index(r"\caption{Hyper-parameters") < tex.index(r"\begin{tabular}"))
check("credits environment", r"\subsubsection{\discintname}" in tex and r"\subsubsection{\ackname}" in tex)
check("thebibliography with 13 bibitems and \\doi", tex.count(r"\bibitem{") == 13 and r"\doi{10.1038/nature14539}" in tex)
check("headings not numbered by hand", r"\section{Introduction}" in tex and r"\subsubsection{Acceptance Test}" in tex)
check("no stray markdown", "**" not in tex and "$$" not in tex)

engine = next((e for e in ("latexmk", "pdflatex", "tectonic") if shutil.which(e)), None)
if engine:
    p = run(BUILD, d / "paper.md", "--latex", d / "latex", "--compile")
    check("latex compiles", p.returncode == 0 and (d / "latex" / "paper.pdf").is_file(), p.stderr[-800:])
else:
    skip("latex compiles", "no latexmk/pdflatex/tectonic on PATH")

# ---------------------------------------------------------------- build errors

d = workdir("errors")
meta, body = split_source(d / "paper.md")
bad_meta = dict(meta)
bad_meta.pop("title")
bad_meta["authors"] = [dict(meta["authors"][0], affiliations=[7])]
write_source(d / "bad.md", bad_meta, body)
p = run(BUILD, d / "bad.md", "--docx", d / "bad.docx", "--latex", d / "badtex")
check("build fails on missing title / bad affiliation number",
      p.returncode == 1 and "title: missing" in p.stderr and "not an affiliation number" in p.stderr, p.stderr)
check("failed build writes nothing", not (d / "bad.docx").exists() and not (d / "badtex").exists())
p = run(BUILD, d / "paper.md", "--docx", d / "paper.docm")
check("build refuses .docm", p.returncode == 1 and ".docx" in p.stderr)
write_source(d / "img.md", meta, body.replace("fig1.png", "missing.png"))
p = run(BUILD, d / "img.md", "--docx", d / "img.docx")
check("build fails on missing image", p.returncode == 1 and not (d / "img.docx").exists(), p.stderr[-300:])

# ---------------------------------------------------------------- checker: example is clean of errors

d = workdir("check")
run(BUILD, d / "paper.md", "--docx", d / "paper.docx")
p = run(CHECK, d / "paper.md", "--docx", d / "paper.docx", "--year", "2026")
check("example: no ERRORs (source + docx)", p.returncode == 0 and "SUMMARY: 0 error(s)" in p.stdout, p.stdout)
check("example: placeholder ORCIDs flagged", "authors.placeholder" in codes(p.stdout))
check("example: reminders present", {"submit.ltp", "template.flavour", "ai.policy"} <= codes(p.stdout))

# ---------------------------------------------------------------- checker: broken sources

meta, body = split_source(d / "paper.md")
broken = dict(meta)
broken["title"] = "a study of drift in boosted models."
broken["abstract"] = "Too short [3]."
broken["keywords"] = ["drift"]
broken["authors"] = [dict(a, corresponding=True) for a in meta["authors"]]
broken["authors"][1]["orcid"] = "1234"
broken["authors"][0]["given"] = "I."
broken["affiliations"] = meta["affiliations"] + ["Unused Institute"]
broken.pop("disclosure")
broken["references"] = meta["references"] + ["Петренко, І.: Назва. Журнал 1, 1–2 (2020)", "[14] Numbered, A.: X. Y (2020)"]
broken_body = (body.replace("# Introduction", "# 1 Introduction")
               .replace("# Related Work", "# related work")
               .replace("[1]", "[2]", 1)
               .replace('alt="Line chart: validation error falls steeply during the first 100 boosting rounds and then flattens."', "")
               + "\n##### Too Deep\n\nText.\n")
write_source(d / "broken.md", broken, broken_body)
p = run(CHECK, d / "broken.md", "--no-reminders", "--year", "2026")
errs, warns = level_codes(p.stdout, "ERROR"), level_codes(p.stdout, "WARN")
check("broken: exits 1", p.returncode == 1)
for code in ("title.period", "abstract.words", "authors.corresponding", "authors.orcid", "affiliations.unused",
             "credits.disclosure", "refs.latin", "refs.numbering", "cite.unused", "cite.order", "headings.numbered",
             "headings.depth"):
    check(f"broken: ERROR {code}", code in errs, p.stdout)
for code in ("title.case", "keywords.count", "keywords.case", "abstract.citations", "authors.initials",
             "figure.alt", "headings.case"):
    check(f"broken: WARN {code}", code in warns, p.stdout)

nofig = dict(meta)
write_source(d / "nofig.md", nofig, body.replace("Figure 1 shows", "This shows").replace("Table 1", "the table"))
p = run(CHECK, d / "nofig.md", "--no-reminders")
check("unreferenced figure/table -> ERROR figure.ref / table.ref",
      {"figure.ref", "table.ref"} <= level_codes(p.stdout, "ERROR"), p.stdout)

alpha = dict(meta, citation_order="alphabetical")
write_source(d / "alpha.md", alpha, body)
p = run(CHECK, d / "alpha.md", "--no-reminders")
check("alphabetical order declared but list unsorted -> cite.order", "cite.order" in level_codes(p.stdout, "ERROR"))

# ---------------------------------------------------------------- renumber_refs

shutil.copy(d / "alpha.md", d / "renum.md")
p = run(RENUM, d / "renum.md", "--order", "alphabetical")
check("renumber alphabetical runs", p.returncode == 0 and (d / "renum.md.bak").is_file(), p.stderr)
m2, b2 = split_source(d / "renum.md")
check("renumbered list is sorted", [lc.sort_key(r) for r in m2["references"]] == sorted(lc.sort_key(r) for r in m2["references"]))
check("renumber keeps every reference", sorted(m2["references"]) == sorted(meta["references"]))
first_ref = m2["references"][0]
check("citations follow the new numbers",
      first_ref.startswith("Borisov") and "[1, 5, 11]" in b2, b2[:600])
p = run(CHECK, d / "renum.md", "--no-reminders")
check("renumbered source passes the citation checks", not ({"cite.order", "cite.unused", "cite.range"} & codes(p.stdout)), p.stdout)
p = run(RENUM, d / "renum.md", "--order", "citation")
p = run(CHECK, d / "renum.md", "--no-reminders")
check("renumber back to citation order passes", "cite.order" not in codes(p.stdout) and "citation_order: citation" in (d / "renum.md").read_text())
check("format_citation compresses ranges", lc.format_citation([7, 3, 4, 5, 9]) == "[3–5, 7, 9]")

# ---------------------------------------------------------------- import

d = workdir("import")
run(BUILD, d / "paper.md", "--docx", d / "paper.docx", "--latex", d / "latex")
p = run(IMPORT, d / "paper.docx", "--out", d / "imp" / "paper.md")
check("import docx runs", p.returncode == 0, p.stderr)
mi, bi = split_source(d / "imp" / "paper.md")
check("import docx: 13 references, abstract, 5 keywords",
      len(mi["references"]) == 13 and mi["abstract"].startswith("Gradient-boosted") and len(mi["keywords"]) == 5,
      f"{len(mi['references'])} {mi['abstract'][:30]} {mi['keywords']}")
check("import docx: heading numbers stripped", "# Introduction" in bi and "# 1" not in bi)
p = run(IMPORT, d / "paper.docx", "--out", d / "imp" / "paper.md")
check("import refuses to overwrite", p.returncode == 1 and "--force" in p.stderr)
p = run(IMPORT, d / "latex" / "paper.tex", "--out", d / "imp2" / "paper.md")
mt, bt = split_source(d / "imp2" / "paper.md")
check("import tex: title, refs with DOI, keywords, credits",
      mt["title"].startswith("Drift-Aware") and len(mt["references"]) == 13
      and "https://doi.org/10.1038/nature14539" in mt["references"][0] and len(mt["keywords"]) == 5
      and mt["disclosure"].startswith("The authors have no"), str(mt)[:500])
check("import tex: \\cite mapped back to numbers", "[5, 6, 7]" in bt and "@ref" not in bt)
p = run(IMPORT, d / "paper.docx", "--out", d / "x.md", "--force")
p = run(CHECK, d / "imp" / "paper.md", "--no-reminders")
check("imported skeleton reports TODOs", "meta.todo" in codes(p.stdout))

# ---------------------------------------------------------------- render + pdf checks

if shutil.which("soffice") or Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists():
    d = workdir("render")
    run(BUILD, d / "paper.md", "--docx", d / "paper.docx")
    p = run(PREVIEW, d / "paper.docx", "--out", d / "preview")
    pdf = d / "preview" / "paper.pdf"
    check("render produces pdf + pages + contact sheet",
          p.returncode == 0 and pdf.is_file() and (d / "preview" / "contact-sheet.png").is_file(), p.stderr[-500:])
    if pdf.is_file():
        p = run(CHECK, d / "paper.md", "--pdf", pdf, "--no-reminders")
        check("pdf check: page count reported, no type3", "pages" in codes(p.stdout) and "pdf.type3" not in codes(p.stdout),
              p.stdout)
        tight = dict(split_source(d / "paper.md")[0])
        tight["conference"] = {"name": "X", "max_pages": 2}
        write_source(d / "tight.md", tight, split_source(d / "paper.md")[1])
        p = run(CHECK, d / "tight.md", "--pdf", pdf, "--no-reminders")
        check("pdf over conference limit -> ERROR pages", "pages" in level_codes(p.stdout, "ERROR"), p.stdout)
        p = run(PREVIEW, pdf, "--out", d / "preview2")
        check("render accepts a PDF directly", p.returncode == 0 and (d / "preview2" / "page-01.png").is_file())
else:
    skip("render + pdf checks", "LibreOffice not installed")

# ---------------------------------------------------------------- real published example

REAL = SKILL_DIR / "assets" / "examples" / "en-ils-simulator-review"
d = WORK / "real"
shutil.copytree(REAL, d)
p = run(BUILD, d / "paper.md", "--docx", d / "paper.docx", "--latex", d / "latex", "--verbose")
check("real example builds (docx + latex)", p.returncode == 0 and (d / "latex" / "paper.tex").is_file(), p.stderr[-800:])
check("real example: 6 sections, 1 figure, 1 table, 24 references",
      "sections=6, figures=1, tables=1, equations=0, references=24" in p.stderr, p.stderr[-300:])
p = run(CHECK, d / "paper.md", "--docx", d / "paper.docx", "--year", "2026")
# The chapter prints no Disclosure of Interests and cites [8] before [7]; these three findings are
# documented in the example's README and are the only ERRORs allowed.
real_errors = level_codes(p.stdout, "ERROR")
check("real example: no ERRORs beyond the documented published-chapter findings",
      real_errors <= {"credits.disclosure", "docx.disclosure", "cite.order"}, p.stdout)
check("real example: no placeholder or TODO findings", not {"authors.placeholder", "todo"} & codes(p.stdout), p.stdout)

# ---------------------------------------------------------------- summary

failed = [r for r in RESULTS if not r[1]]
print(f"SUMMARY: {len(RESULTS) - len(failed)} passed, {len(failed)} failed, {len(SKIPPED)} skipped")
sys.exit(1 if failed else 0)
